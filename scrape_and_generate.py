#scrape_and_generate.py
import re, sys, time, json, threading
from pathlib import Path
from datetime import date
import yaml
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
from io import BytesIO
from docx import Document
from docx.oxml import parse_xml
from datetime import datetime
import os
import shutil
import subprocess
import socket
import urllib.request
from docx.table import _Cell, Table
from docx.oxml.ns import qn
import html
from xml.sax.saxutils import escape as _xml_escape
import logging
from datetime import datetime
def resolve_user_data_dir(cfg: dict) -> Path:
    raw = (cfg.get("user_data_dir") or "").strip()
    if not raw:
        raw = r"%LOCALAPPDATA%\CustomerLetterGenerator\gch_browser_profile"
    path = Path(os.path.expandvars(os.path.expanduser(raw))).resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path
def _kill_browser_processes_for_profile_repair():
    if not sys.platform.startswith("win"):
        return
    run_kwargs = {
        "capture_output": True,
        "text": True,
        "timeout": 20,
        "check": False,
    }
    if hasattr(subprocess, "CREATE_NO_WINDOW"):
        run_kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
    for exe_name in ("msedge.exe", "msedgewebview2.exe", "chrome.exe"):
        try:
            subprocess.run(["taskkill", "/F", "/IM", exe_name], **run_kwargs)
            try:
                log(f"[browser-repair] Closed process if running: {exe_name}")
            except Exception:
                pass
        except Exception as exc:
            try:
                log(f"[browser-repair] taskkill {exe_name} failed: {exc}")
            except Exception:
                pass
def _remove_profile_dir(profile_dir: Path):
    try:
        if profile_dir.exists():
            shutil.rmtree(profile_dir, ignore_errors=True)
        if profile_dir.exists() and any(profile_dir.iterdir()):
            backup = profile_dir.with_name(
                profile_dir.name + "_bad_" + datetime.now().strftime("%Y%m%d_%H%M%S")
            )
            try:
                profile_dir.rename(backup)
                log(f"[browser-repair] Renamed locked/damaged profile to: {backup}")
            except Exception as exc:
                log(f"[browser-repair] Could not rename damaged profile {profile_dir}: {exc}")
        profile_dir.mkdir(parents=True, exist_ok=True)
        log(f"[browser-repair] Clean profile ready: {profile_dir}")
    except Exception as exc:
        log(f"[browser-repair] Failed to clean profile {profile_dir}: {exc}")
        raise
def _reset_browser_profiles(main_profile_dir: Path):
    _kill_browser_processes_for_profile_repair()
    time.sleep(1)
    app_dir = main_profile_dir.parent
    profile_dirs = [
        main_profile_dir,
        app_dir / "gch_browser_profile_chromium",
    ]
    for profile_dir in profile_dirs:
        _remove_profile_dir(profile_dir)
def _edge_executable_candidates():
    paths = [
        Path(os.environ.get("ProgramFiles(x86)", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
        Path(os.environ.get("ProgramFiles", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
    ]
    seen = set()
    for p in paths:
        if not str(p):
            continue
        key = str(p).lower()
        if key not in seen:
            seen.add(key)
            if p.exists():
                yield p
    found = shutil.which("msedge.exe") or shutil.which("msedge")
    if found:
        p = Path(found)
        key = str(p).lower()
        if key not in seen and p.exists():
            yield p
def _find_edge_executable() -> Path:
    for p in _edge_executable_candidates():
        return p
    raise RuntimeError("Could not find Microsoft Edge executable (msedge.exe).")
def _find_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))
        return int(s.getsockname()[1])
def _wait_for_cdp(port: int, timeout_seconds: int = 30):
    url = f"http://127.0.0.1:{port}/json/version"
    deadline = time.monotonic() + timeout_seconds
    last_error = None
    while time.monotonic() < deadline:
        try:
            with urllib.request.urlopen(url, timeout=2) as resp:
                if resp.status == 200:
                    return
        except Exception as exc:
            last_error = exc
            time.sleep(0.5)
    raise RuntimeError(f"Timed out waiting for Edge CDP endpoint on port {port}: {last_error}")
def _collect_child_process_ids(parent_pid: int) -> set[int]:
    if not sys.platform.startswith("win"):
        return {parent_pid}
    try:
        import ctypes
        from ctypes import wintypes
        kernel32 = ctypes.windll.kernel32
        TH32CS_SNAPPROCESS = 0x00000002
        INVALID_HANDLE_VALUE = ctypes.c_void_p(-1).value
        class PROCESSENTRY32(ctypes.Structure):
            _fields_ = [
                ("dwSize", wintypes.DWORD),
                ("cntUsage", wintypes.DWORD),
                ("th32ProcessID", wintypes.DWORD),
                ("th32DefaultHeapID", ctypes.POINTER(wintypes.ULONG)),
                ("th32ModuleID", wintypes.DWORD),
                ("cntThreads", wintypes.DWORD),
                ("th32ParentProcessID", wintypes.DWORD),
                ("pcPriClassBase", wintypes.LONG),
                ("dwFlags", wintypes.DWORD),
                ("szExeFile", ctypes.c_char * 260),
            ]
        snapshot = kernel32.CreateToolhelp32Snapshot(TH32CS_SNAPPROCESS, 0)
        if snapshot == INVALID_HANDLE_VALUE:
            return {parent_pid}
        try:
            entry = PROCESSENTRY32()
            entry.dwSize = ctypes.sizeof(PROCESSENTRY32)
            children_by_parent = {}
            if not kernel32.Process32First(snapshot, ctypes.byref(entry)):
                return {parent_pid}
            while True:
                pid = int(entry.th32ProcessID)
                ppid = int(entry.th32ParentProcessID)
                children_by_parent.setdefault(ppid, set()).add(pid)
                if not kernel32.Process32Next(snapshot, ctypes.byref(entry)):
                    break
        finally:
            kernel32.CloseHandle(snapshot)
        all_pids = {parent_pid}
        pending = [parent_pid]
        while pending:
            current = pending.pop()
            for child_pid in children_by_parent.get(current, set()):
                if child_pid not in all_pids:
                    all_pids.add(child_pid)
                    pending.append(child_pid)
        return all_pids
    except Exception:
        return {parent_pid}
def _set_process_windows_enabled(process, enabled: bool):
    if not sys.platform.startswith("win") or process is None:
        return
    try:
        import ctypes
        from ctypes import wintypes
        user32 = ctypes.windll.user32
        target_pids = _collect_child_process_ids(int(process.pid))
        windows = []
        EnumWindowsProc = ctypes.WINFUNCTYPE(ctypes.c_bool, wintypes.HWND, wintypes.LPARAM)
        def callback(hwnd, _lparam):
            pid = wintypes.DWORD()
            user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
            if int(pid.value) in target_pids and user32.IsWindowVisible(hwnd):
                windows.append(hwnd)
            return True
        user32.EnumWindows(EnumWindowsProc(callback), 0)
        for hwnd in windows:
            user32.EnableWindow(hwnd, bool(enabled))
    except Exception as exc:
        try:
            log(f"[browser] Could not {'enable' if enabled else 'disable'} Edge window input: {exc}")
        except Exception:
            pass
class _BrowserInputLock:
    def __init__(self, process, interval_seconds: float = 0.75):
        self._process = process
        self._interval_seconds = interval_seconds
        self._stop_event = threading.Event()
        self._thread = None
    def start(self):
        if not sys.platform.startswith("win") or self._process is None:
            return
        _set_process_windows_enabled(self._process, False)
        self._thread = threading.Thread(target=self._run, name="EdgeInputLock", daemon=True)
        self._thread.start()
    def _run(self):
        while not self._stop_event.wait(self._interval_seconds):
            if self._process.poll() is not None:
                return
            _set_process_windows_enabled(self._process, False)
    def stop(self):
        self._stop_event.set()
        if self._thread and self._thread.is_alive():
            self._thread.join(timeout=2)
        _set_process_windows_enabled(self._process, True)
class _CdpBrowserContext:
    def __init__(self, browser, context, process=None, input_lock=None):
        self._browser = browser
        self._context = context
        self._process = process
        self._input_lock = input_lock
        self._input_released = False
    def __getattr__(self, name):
        return getattr(self._context, name)
    def release_input(self):
        if self._input_released:
            return
        if self._input_lock:
            self._input_lock.stop()
        else:
            _set_process_windows_enabled(self._process, True)
        self._input_released = True
    def close(self):
        try:
            self.release_input()
            self._browser.close()
        finally:
            if self._process and self._process.poll() is None:
                try:
                    self._process.terminate()
                except Exception:
                    pass
def _launch_edge_over_cdp(playwright, cfg: dict, profile_dir: Path):
    edge_exe = _find_edge_executable()
    port = int(cfg.get("edge_debugging_port") or 0) or _find_free_port()
    args = [
        str(edge_exe),
        f"--remote-debugging-port={port}",
        f"--user-data-dir={profile_dir}",
        "--no-first-run",
        "--new-window",
        "about:blank",
    ]
    log(f"[browser] Starting Edge over CDP: {edge_exe}")
    log(f"[browser] Edge CDP profile: {profile_dir}")
    log(f"[browser] Edge CDP port: {port}")
    popen_kwargs = {}
    if hasattr(subprocess, "CREATE_NO_WINDOW"):
        popen_kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
    proc = subprocess.Popen(args, **popen_kwargs)
    try:
        _wait_for_cdp(port, timeout_seconds=45)
        input_lock = _BrowserInputLock(proc)
        input_lock.start()
        log("[browser] Edge window input disabled until scraping completes.")
        browser = playwright.chromium.connect_over_cdp(f"http://127.0.0.1:{port}")
        context = browser.contexts[0] if browser.contexts else browser.new_context()
        wrapped_context = _CdpBrowserContext(browser, context, proc, input_lock)
        page = context.pages[0] if context.pages else context.new_page()
        return wrapped_context, page
    except Exception:
        try:
            input_lock.stop()
        except UnboundLocalError:
            _set_process_windows_enabled(proc, True)
        if proc.poll() is None:
            try:
                proc.terminate()
            except Exception:
                pass
        raise
def _new_launch_kwargs(cfg: dict, include_channel: bool = True) -> dict:
    launch_kwargs = {
        "headless": bool(cfg.get("headless", False)),
        "viewport": None,
        "accept_downloads": True,
        "args": [
            "--start-maximized",
        ],
    }
    browser_channel = (cfg.get("browser_channel") or "").strip()
    if include_channel and browser_channel:
        launch_kwargs["channel"] = browser_channel
    return launch_kwargs
def _launch_persistent_context(playwright, profile_dir: Path, launch_kwargs: dict):
    return playwright.chromium.launch_persistent_context(
        user_data_dir=str(profile_dir),
        **launch_kwargs,
    )
def launch_gch_context(playwright, cfg: dict):
    user_data_dir = resolve_user_data_dir(cfg)
    browser_channel = (cfg.get("browser_channel") or "").strip().lower()
    if browser_channel in {"msedge", "edge"}:
        try:
            return _launch_edge_over_cdp(playwright, cfg, user_data_dir)
        except Exception as first_error:
            log(f"[browser] Edge CDP launch failed: {first_error}")
            log("[browser] Resetting GCH Edge profile and retrying Edge CDP once.")
            _reset_browser_profiles(user_data_dir)
            try:
                return _launch_edge_over_cdp(playwright, cfg, user_data_dir)
            except Exception as second_error:
                log(f"[browser] Edge CDP launch after reset failed: {second_error}")
                raise RuntimeError(
                    "Microsoft Edge could not be opened for GCH automation. The app already "
                    "closed Edge/WebView2, reset the GCH Edge profile, removed the old Chromium "
                    "fallback profile, and retried automatically.\n\n"
                    "The app did not fall back to Chromium because Chromium can trigger "
                    "Microsoft/Medtronic MFA verification problems.\n\n"
                    "Please reopen the app and try again. If this repeats, Edge may be blocked "
                    "from remote debugging by a corporate policy on this computer."
                ) from second_error
    launch_kwargs = _new_launch_kwargs(cfg, include_channel=True)
    context = _launch_persistent_context(playwright, user_data_dir, launch_kwargs)
    page = context.pages[0] if context.pages else context.new_page()
    return context, page
def setup_logging():
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    log_filename = f"customer_letter_debug_{timestamp}.log"
    logger = logging.getLogger()
    logger.setLevel(logging.DEBUG)
    logger.handlers = []
    file_handler = logging.FileHandler(log_filename, encoding='utf-8')
    file_handler.setLevel(logging.DEBUG)
    file_formatter = logging.Formatter('[%(asctime)s] %(message)s', datefmt='%Y-%m-%d %H:%M:%S')
    file_handler.setFormatter(file_formatter)
    logger.addHandler(file_handler)
    console_handler = logging.StreamHandler()
    console_handler.setLevel(logging.DEBUG)
    console_formatter = logging.Formatter('[%(asctime)s] %(message)s', datefmt='%Y-%m-%d %H:%M:%S')
    console_handler.setFormatter(console_formatter)
    logger.addHandler(console_handler)
    return log_filename
def log(msg):
    logging.info(msg)
def ts():
    return time.strftime("%Y-%m-%d %H:%M:%S")
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "0")
ENDS_WITH_PRODUCT     = "substring(@id, string-length(@id) - string-length('-Product') + 1) = '-Product'"
ENDS_WITH_DESCRIPTION = "substring(@id, string-length(@id) - string-length('-Description') + 1) = '-Description'"
_WT_RX = re.compile(r'(<w:t\b[^>]*>)(.*?)(</w:t>)', re.S | re.I)
def get_user_data_dir(app_name="CustomerLetterGenerator"):
    base = os.getenv("LOCALAPPDATA") or os.getenv("APPDATA") or str(Path.home())
    d = Path(base) / app_name / "chrome-profile"
    d.mkdir(parents=True, exist_ok=True)
    return str(d)
def _xml_convert_newlines_to_br(xml: str) -> str:
    def repl(m):
        open_tag, text, close_tag = m.group(1), m.group(2), m.group(3)
        if 'xml:space' not in open_tag:
            open_tag = open_tag[:-1] + ' xml:space="preserve">'
        if '\n' not in text:
            return open_tag + text + close_tag
        parts = text.split('\n')
        return open_tag + ('</w:t><w:br/><w:t xml:space="preserve">'.join(parts)) + close_tag
    return _WT_RX.sub(repl, xml)
def _convert_bullets_to_indented_paras(xml: str) -> str:
    BULLET = '\u2022'
    def find_rpr_before(xml_str, pos):
        chunk = xml_str[max(0, pos-3000):pos]
        matches = list(re.finditer(r'<w:rPr>(.*?)</w:rPr>', chunk, re.S))
        if matches:
            return '<w:rPr>' + matches[-1].group(1) + '</w:rPr>'
        return ''
    def find_ppr_before(xml_str, pos):
        chunk = xml_str[max(0, pos-3000):pos]
        matches = list(re.finditer(r'<w:pPr>(.*?)</w:pPr>', chunk, re.S))
        if matches:
            return '<w:pPr>' + matches[-1].group(1) + '</w:pPr>'
        return ''
    bullet_pattern = re.compile(
        r'(</w:t>)\s*(<w:br/>)\s*(<w:t[^>]*>)\s*' + re.escape(BULLET) + r'[ \t]*',
        re.UNICODE
    )
    matches = list(bullet_pattern.finditer(xml))
    for m in reversed(matches):
        rpr = find_rpr_before(xml, m.start())
        replacement = (
            '</w:t></w:r></w:p>'
            '<w:p><w:pPr>'
            '<w:ind w:left="720" w:hanging="360"/>'
            '</w:pPr>'
            f'<w:r>{rpr}<w:t xml:space="preserve">{BULLET}\t'
        )
        xml = xml[:m.start()] + replacement + xml[m.end():]
    br_pattern = re.compile(
        r'(</w:t>)\s*(<w:br/>)\s*(<w:t[^>]*>)',
    )
    def is_in_indented_para(xml_str, pos):
        chunk = xml_str[max(0, pos-2000):pos]
        last_p = max(chunk.rfind('<w:p>'), chunk.rfind('<w:p><w:pPr>'))
        if last_p == -1:
            return False
        return 'w:ind' in chunk[last_p:]
    matches2 = list(br_pattern.finditer(xml))
    for m in reversed(matches2):
        if is_in_indented_para(xml, m.start()):
            rpr = find_rpr_before(xml, m.start())
            after = xml[m.end():m.end()+10]
            if after.lstrip().startswith(BULLET):
                continue
            original_ppr = find_ppr_before(xml, m.start())
            clean_ppr = re.sub(r'<w:ind[^/]*/>', '', original_ppr)
            if clean_ppr == '<w:pPr></w:pPr>':
                clean_ppr = ''
            replacement = (
                '</w:t></w:r></w:p>'
                f'<w:p>{clean_ppr}'
                f'<w:r>{rpr}<w:t xml:space="preserve">'
            )
            xml = xml[:m.start()] + replacement + xml[m.end():]
    return xml
def _norm(s: str) -> str:
    return re.sub(r'\s+', ' ', (s or '').replace('\xa0',' ')).strip().lower()
def _cell_txt(cell: _Cell) -> str:
    return _norm(cell.text)
def _delete_row(table: Table, row_idx: int):
    tr = table.rows[row_idx]._tr
    table._tbl.remove(tr)
def _clean_sn(sn: str) -> str:
    s = (sn or "").strip()
    s = re.sub(r'^\s*(sn|s/?n)\s*[:#-]?\s*', '', s, flags=re.I).strip()
    if s.lower() in {"", "na", "n/a", "none", "-"}:
        return ""
    return s
def _clean_lot(lot: str) -> str:
    s = (lot or "").strip()
    s = re.sub(r'^\s*(ln|lot)\s*[:#-]?\s*', '', s, flags=re.I).strip()
    if s.lower() in {"", "lot", "na", "n/a", "none", "-"}:
        return ""
    return s
def _join_serial_lot(p) -> str:
    sn  = _clean_sn(p.get('sn'))
    lot = _clean_lot(p.get('lot'))
    if sn and lot:
        return f"SN: {sn} / LN: {lot}"
    elif sn:
        return f"SN: {sn}"
    elif lot:
        return f"LN: {lot}"
    else:
        return ""
def _row_text(row) -> str:
    return " ".join(_cell_txt(c) for c in row.cells)
def _looks_like_products_header_row(row) -> bool:
    cells = row.cells
    if len(cells) < 3:
        return False
    a, b, c = _cell_txt(cells[0]), _cell_txt(cells[1]), _cell_txt(cells[2])
    ok0 = a.startswith("product")
    ok1 = ("product" in b and "description" in b)
    ok2 = ("serial" in c and "lot" in c)
    return ok0 and ok1 and ok2
def _find_products_table(doc: Document):
    for tbl in doc.tables:
        for r_idx, row in enumerate(tbl.rows):
            if _looks_like_products_header_row(row):
                return tbl, r_idx
        for r_idx, row in enumerate(tbl.rows):
            t = _row_text(row)
            if "product id2" in t or "product_desc2" in t.replace(" ", "_") or "{{product" in t or "[[product" in t:
                return tbl, 0
    return None, None
def _clear_rows_after(tbl: Table, header_row_idx: int):
    for i in range(len(tbl.rows) - 1, header_row_idx, -1):
        _delete_row(tbl, i)
def _fill_row(cells, p):
    pid = (p.get("id") or p.get("code") or "").strip()
    cells[0].text = pid
    cells[1].text = (p.get("desc") or "").strip()
    cells[2].text = _join_serial_lot(p)
def _update_products_table(doc: Document, products: list):
    tbl, hdr = _find_products_table(doc)
    if not tbl:
        print("[DOCX] Products table not found; skipping dynamic rows.")
        return
    body_indices = list(range(hdr + 1, len(tbl.rows)))
    if not products:
        for i in reversed(body_indices):
            _delete_row(tbl, i)
        return
    if len(products) == 1:
        if body_indices:
            first_idx = body_indices[0]
            _fill_row(tbl.rows[first_idx].cells, products[0])
            for i in reversed(body_indices[1:]):
                _delete_row(tbl, i)
        else:
            row = tbl.add_row()
            _fill_row(row.cells, products[0])
        return
    _clear_rows_after(tbl, hdr)
    for p in products:
        row = tbl.add_row()
        _fill_row(row.cells, p)
def find_first_visible_input(page, primary_selector, fallbacks=None, timeout=15000):
    fallbacks = fallbacks or []
    selectors = [primary_selector] + fallbacks
    contexts = [page] + [fr for fr in page.frames]
    last_exc = None
    for sel in selectors:
        for ctx in contexts:
            try:
                loc = ctx.locator(sel).first
                try:
                    loc.wait_for(state="visible", timeout=timeout)
                    return loc, ctx
                except Exception as e:
                    last_exc = e
                count = loc.count()
                if count:
                    return loc, ctx
            except Exception as e:
                last_exc = e
                continue
    raise PWTimeout(f"Could not find element (visible or attached) for selectors: {selectors}")
def dump_frames_debug(page, basename="debug"):
    info_lines = []
    frames = page.frames
    for idx, fr in enumerate(frames):
        try:
            url = fr.url or ""
            name = fr.name or ""
        except Exception:
            url = ""
            name = ""
        info_lines.append(f"[{idx}] name={name!r} url={url!r}")
        try:
            pass
        except Exception:
            pass
        try:
            candidates = fr.locator("xpath=//input[contains(@class,'th-sif') or contains(@id,'SearchValue') or contains(@tempname,'search')]")
            count = candidates.count()
            info_lines.append(f"  candidates(th-sif/id*SearchValue/tempname*search): {count}")
            for j in range(min(count, 10)):
                el = candidates.nth(j)
                html = el.evaluate("e => e.outerHTML").strip()
                info_lines.append("    " + html.replace("\n"," "))
        except Exception as e:
            info_lines.append(f"  (error collecting candidates: {e})")
    try:
        Path(f"{basename}_frames.txt").write_text("\n".join(info_lines), encoding="utf-8")
    except Exception:
        pass
def soft_click_go(ctx, custom_selector=None):
    candidates = []
    if custom_selector:
        candidates.append(custom_selector)
    candidates += [
        "text=Go",
        "input[value='Go']",
        "button:has-text('Go')",
        "xpath=//a[normalize-space(.)='Go']",
    ]
    for sel in candidates:
        try:
            btn = ctx.locator(sel).first
            if btn.count():
                btn.click()
                return True
        except Exception:
            pass
    return False
def wait_find_in_any_frame(page, selectors, timeout_ms=30000, poll_ms=600):
    import time
    deadline = time.time() + (timeout_ms/1000.0)
    tried = set()
    while time.time() < deadline:
        log(f"[wait] scanning {len(page.frames)} frames for any of: {selectors}")
        frames = page.frames
        for sel in selectors:
            if sel in tried:
                pass
            for fr in frames:
                try:
                    loc = fr.locator(sel).first
                    if loc.count():
                        try:
                            loc.wait_for(state="visible", timeout=poll_ms)
                        except Exception:
                            pass
                        log(f"[wait] found {sel} in frame name={getattr(fr,'name','')} url={getattr(fr,'url','')}")
                        return loc, fr, sel
                except Exception:
                    continue
        time.sleep(poll_ms/1000.0)
    raise PWTimeout(f"Could not find element in any frame for selectors: {selectors}")
import re
from docx import Document
def _norm_key(s: str) -> str:
    return re.sub(r'[^a-z0-9]+', '_', (s or '').strip().lower()).strip('_')
DEFAULT_SIGNATURE_MANAGER_NAME = 'Tracy Landers'
DEFAULT_SIGNATURE_MANAGER_TITLE = 'Sr MDR/Vigilance Manager'
def _build_alias_mapping(mapping: dict) -> dict:
    out = {}
    for k, v in mapping.items():
        nk = _norm_key(k)
        out[nk] = v
        if nk.endswith('_1'):
            out[nk[:-2]] = v
    for i in range(1, 10):
        a_under = f"analysis_{i}"; a_flat = f"analysis{i}"
        if a_under in out and a_flat not in out: out[a_flat] = out[a_under]
        if a_flat  in out and a_under not in out: out[a_under] = out[a_flat]
        iv_under = f"investigation_{i}"; iv_flat = f"investigation{i}"
        if iv_under in out and iv_flat not in out: out[iv_flat] = out[iv_under]
        if iv_flat  in out and iv_under not in out: out[iv_under] = out[iv_flat]
    aliases = {
        'today_date': out.get('todays_date', ''),
        "today_s_date": out.get('todays_date', ''),
        'ir_name': out.get('ir_name', ''),
        'ir_with_address': out.get('ir_with_address', ''),
        'event_date': out.get('event_date', ''),
        'event_description': out.get('event_description', ''),
        'analysis_results_if_present': out.get('analysis_results', ''),
        'investigation_summary': out.get('investigation_summary', ''),
        'signature_manager_name': out.get('signature_manager_name') or DEFAULT_SIGNATURE_MANAGER_NAME,
        'signature_manager_title': out.get('signature_manager_title') or DEFAULT_SIGNATURE_MANAGER_TITLE,
        'product_id': out.get('product_id_1', ''),
        'product_desc': out.get('product_desc_1', ''),
        'lot_serial_number': out.get('serial_or_lot_1', ''),
        'lot_serial_no': out.get('serial_or_lot_1', ''),
        'lot/serial number': out.get('serial_or_lot_1', ''),
        'serial no/lot no': out.get('serial_or_lot_1', ''),
        'serial_no_lot_no': out.get('serial_or_lot_1', ''),
        'pe_number': out.get('complaint_id', ''),
        'pe number': out.get('complaint_id', ''),
        'product_id2': out.get('product_id_2', ''),
        'product id2': out.get('product_id_2', ''),
        'product_desc2': out.get('product_desc_2', ''),
        'product desc2': out.get('product_desc_2', ''),
        'product description2': out.get('product_desc_2', ''),
        'lot_serial_number2': out.get('serial_or_lot_2', ''),
        'lot/serial number2': out.get('serial_or_lot_2', ''),
        'serial no/lot no2': out.get('serial_or_lot_2', ''),
        'serial_no_lot_no2': out.get('serial_or_lot_2', ''),
        'ex ref': out.get('ex_ref', ''),
        'external reference': out.get('ex_ref', ''),
        'external_reference': out.get('ex_ref', ''),
    }
    if out.get('rb_reference'):
        aliases.update({
            'rb reference': out['rb_reference'],
            'rb_reference': out['rb_reference'],
            'ref number':   out['rb_reference'],
            'ref_number':   out['rb_reference'],
        })
    out.update({k: v for k, v in aliases.items() if v})
    return out
def _split_tolerant(label: str) -> str:
    gap = r'(?:\s|<[^>]*?>)*?'
    toks = [t for t in re.split(r'[^A-Za-z0-9]+', (label or '').strip()) if t]
    if not toks: return ''
    sep = r'(?:\s|[/\-\._]|<[^>]*?>)*?'
    parts = []
    for i, tok in enumerate(toks):
        parts.append(gap.join(re.escape(c) for c in tok))
        if i < len(toks) - 1:
            parts.append(sep)
    return gap + ''.join(parts) + gap
def _patterns_for_key(human_label: str):
    gap = r'(?:\s|<[^>]*?>)*?'
    inner = _split_tolerant(human_label)
    flags = re.I | re.S
    return (
        re.compile(r'\[' + gap + r'\[' + r'\s*' + inner + r'\s*' + r'\]' + gap + r'\]', flags),
        re.compile(r'\{' + gap + r'\{' + r'\s*' + inner + r'\s*' + r'\}' + gap + r'\}', flags),
    )
from xml.sax.saxutils import escape as _xml_escape
def _xml_replace_all(xml: str, mapping: dict) -> str:
    fast = re.compile(r'(\{\{|\[\[)\s*(.*?)\s*(\}\}|\]\])', re.I | re.S)
    def _escape_value(v) -> str:
        s = str(v)
        return _xml_escape(s, {'"': '&quot;', "'": '&apos;'})
    def _quick(m):
        k = _norm_key(m.group(2))
        if k in mapping:
            v = mapping[k]
            if v is not None and v != '':
                return _escape_value(v)
        return m.group(0)
    xml_new = fast.sub(_quick, xml)
    keys_seen = set()
    for raw_key, value in mapping.items():
        if value in (None, ''):
            continue
        value_str = _escape_value(value)
        for label in {raw_key, raw_key.replace('_', ' ')}:
            if label in keys_seen:
                continue
            keys_seen.add(label)
            pat_sq, pat_cu = _patterns_for_key(label)
            xml_new = pat_sq.sub(lambda _m, vs=value_str: vs, xml_new)
            xml_new = pat_cu.sub(lambda _m, vs=value_str: vs, xml_new)
    return xml_new
_PLACEHOLDER_FINDER = re.compile(
    r'(?P<open>\{\{|\[\[)'
    r'(?:\s|<[^>]*?>)*?'
    r'(?P<label>[^}\]]{1,120}?)'
    r'(?:\s|<[^>]*?>)*?'
    r'(?P<close>\}\}|\]\])',
    re.I | re.S
)
def _remove_report_number_block_docx(doc: Document, report_number: str):
    if report_number:
        return   
    def _norm_txt(s):
        return re.sub(r'\s+', ' ', (s or '').replace('\xa0',' ')).strip().lower()
    def _is_report_line_text(t: str) -> bool:
        t = _norm_txt(t)
        if not t:
            return False
        if re.search(r'\bmpxr\s*report\b',t):
            return True
        return False
    for tbl in doc.tables:
        rows_to_del = []
        for i, row in enumerate(tbl.rows):
            row_txt = " ".join(_norm_txt(c.text) for c in row.cells)
            if _is_report_line_text(row_txt):
                rows_to_del.append(i)
        for i in reversed(rows_to_del):
            tr = tbl.rows[i]._tr
            tbl._tbl.remove(tr)
    for p in list(doc.paragraphs):
        txt = p.text or ""
        if _is_report_line_text(txt) or ("{{" in txt.lower() and "report" in txt.lower()):
            p._element.getparent().remove(p._element)
def _remove_rb_reference_block_docx(doc: Document, rb_value: str):
    if rb_value:
        return
    def _norm_txt(s):
        return re.sub(r'\s+', ' ', (s or '').replace('\xa0',' ')).strip().lower()
    def _is_rb_line_text(t: str) -> bool:
        t = _norm_txt(t)
        if not t:
            return False
        if re.search(r'\brb\s*ref(erence)?\b', t):
            return True
        return False
    for tbl in doc.tables:
        rows_to_del = []
        for i, row in enumerate(tbl.rows):
            row_txt = " ".join(_norm_txt(c.text) for c in row.cells)
            if _is_rb_line_text(row_txt):
                rows_to_del.append(i)
        for i in reversed(rows_to_del):
            tr = tbl.rows[i]._tr
            tbl._tbl.remove(tr)
    for p in list(doc.paragraphs):
        txt = p.text or ""
        if _is_rb_line_text(txt) or ("{{" in txt.lower() and "ref" in txt.lower()):
            p._element.getparent().remove(p._element)
def _remove_event_date_block_docx(doc: Document, event_date: str):
    if event_date:
        return
    def _norm_txt(s):
        return re.sub(r'\s+', ' ', (s or '').replace('\xa0',' ')).strip().lower()
    def _is_event_date_line_text(t: str) -> bool:
        t = _norm_txt(t)
        if not t:
            return False
        if re.search(r'\bevent\s*date\b', t):
            return True
        if re.search(r'\bdate\s*of\s*event\b', t):
            return True
        return False
    for tbl in doc.tables:
        rows_to_del = []
        for i, row in enumerate(tbl.rows):
            row_txt = " ".join(_norm_txt(c.text) for c in row.cells)
            if _is_event_date_line_text(row_txt):
                rows_to_del.append(i)
        for i in reversed(rows_to_del):
            tr = tbl.rows[i]._tr
            tbl._tbl.remove(tr)
    for p in list(doc.paragraphs):
        txt = p.text or ""
        if _is_event_date_line_text(txt):
            p._element.getparent().remove(p._element)
        else:
            low = txt.lower()
            if "{{" in low or "[[" in low:
                if "event_date" in low or ("event" in low and "date" in low):
                    p._element.getparent().remove(p._element)
def replace_everywhere(doc: Document, mapping: dict):
    resolved = _build_alias_mapping(mapping)
    for part in doc.part.package.parts:
        if 'xml' not in getattr(part, 'content_type', ''):
            continue
        try:
            xml = part._element.xml
        except Exception:
            try:
                xml = part.blob.decode('utf-8', errors='ignore')
            except Exception:
                continue
        ph = _list_placeholders(xml)
        if ph:
            print("[DOCX] placeholders detected in part:", getattr(part, 'partname', '<?>'))
            for s in sorted(set(ph)):
                print("   -", s)
        new_xml = _xml_replace_all(xml, resolved)
        plural = (mapping.get('_product_count') or 0) > 1
        new_xml = _apply_plural_s(new_xml, plural)
        new_xml = _xml_convert_newlines_to_br(new_xml)
        new_xml = _convert_bullets_to_indented_paras(new_xml)
        if new_xml != xml:
            print("[DOCX] replacements applied in", getattr(part, 'partname', '<?>'))
            try:
                part._element = parse_xml(new_xml)
            except Exception as e:
                print("[DOCX] parse_xml failed for", getattr(part, 'partname', '<?>'), ":", e)
DOCX_OUTPUT_FONT = "Avenir Next LT Pro"
def _set_run_font(run, font_name: str = DOCX_OUTPUT_FONT):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = parse_xml('<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), font_name)
def _apply_output_font(doc: Document, font_name: str = DOCX_OUTPUT_FONT):
    def apply_paragraphs(paragraphs):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, font_name)
    apply_paragraphs(doc.paragraphs)
    for section in doc.sections:
        apply_paragraphs(section.header.paragraphs)
        apply_paragraphs(section.footer.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                apply_paragraphs(cell.paragraphs)
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            apply_paragraphs(nested_cell.paragraphs)
def _iter_document_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for paragraph in section.footer.paragraphs:
            yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            for paragraph in nested_cell.paragraphs:
                                yield paragraph
def _replace_static_signature_block(doc: Document, mapping: dict):
    resolved = _build_alias_mapping(mapping)
    signature_name = resolved.get('signature_manager_name') or DEFAULT_SIGNATURE_MANAGER_NAME
    signature_title = resolved.get('signature_manager_title') or DEFAULT_SIGNATURE_MANAGER_TITLE
    def set_paragraph_text(paragraph, text):
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ''
        else:
            paragraph.add_run(text)
    for paragraph in _iter_document_paragraphs(doc):
        text = (paragraph.text or '').strip()
        if text == DEFAULT_SIGNATURE_MANAGER_NAME:
            set_paragraph_text(paragraph, signature_name)
        elif text.startswith(DEFAULT_SIGNATURE_MANAGER_TITLE) and 'Customer Quality' in text:
            set_paragraph_text(paragraph, f'{signature_title} | Customer Quality')
def fill_docx(template_path, out_path, mapping, products=None):
    doc = Document(template_path)
    if products is not None:
        _ensure_second_table_product_blocks(doc, len(products))
    replace_everywhere(doc, mapping)
    _replace_static_signature_block(doc, mapping)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    doc = Document(buf)
    _remove_rb_reference_block_docx(doc, (mapping.get('rb_reference') or '').strip())
    _remove_event_date_block_docx(doc, (mapping.get('event_date') or '').strip())
    _remove_report_number_block_docx(doc, (mapping.get('report_number') or '').strip())
    if products is not None:
        _update_products_table(doc, products)
    _apply_output_font(doc)
    doc.save(out_path)
def find_app_frame(page, frame_name_regex=None, url_regex=None):
    if url_regex:
        for fr in page.frames:
            if re.search(url_regex, fr.url or "", re.I):
                return fr
    if frame_name_regex:
        for fr in page.frames:
            if re.search(frame_name_regex, (fr.name or ""), re.I):
                return fr
    return page.main_frame
def clean(s):
    return re.sub(r'\s{2,}', ' ', (s or '').replace('\xa0', ' ')).strip()
def get_by_label(frame, labels):
    if isinstance(labels, str):
        labels = [labels]
    for label in labels:
        loc = frame.locator(f"xpath=//*[normalize-space(text())='{label}']")
        if loc.count() > 0:
            try:
                el = loc.first
                sib = el.locator("xpath=following-sibling::*[1]")
                if sib.count():
                    return clean(sib.first.inner_text())
            except Exception:
                pass
            try:
                for_attr = loc.first.get_attribute('for')
                if for_attr:
                    inp = frame.locator(f"#{for_attr}")
                    if inp.count():
                        val = (inp.first.input_value()
                            if 'input' in (inp.first.evaluate('e => e.tagName') or '').lower()
                            else inp.first.inner_text())
                        return clean(val)
            except Exception:
                pass
            try:
                nxt = loc.first.locator("xpath=following::*[1]")
                if nxt.count():
                    return clean(nxt.first.inner_text())
            except Exception:
                pass
        loc2 = frame.locator(f"xpath=//*[contains(normalize-space(.), '{label}')]")
        if loc2.count():
            try:
                sib = loc2.first.locator("xpath=following-sibling::*[1]")
                if sib.count():
                    return clean(sib.first.inner_text())
            except Exception:
                pass
    return ""
def get_grid_row_col(frame, grid_hint_xpath, row_index, col_index):
    try:
        grid = frame.locator(grid_hint_xpath).first
        if grid.count() == 0:
            return ""
        rows = grid.locator("xpath=.//tr")
        row = rows.nth(row_index)
        cells = row.locator("xpath=.//th|.//td")
        cell = cells.nth(col_index - 1)
        return clean(cell.inner_text())
    except Exception:
        return ""
def click_partners_tab(page, frame):
    sels = [
        "xpath=//a[contains(@class,'ui-tabs-anchor')][normalize-space(.)='Partners']",
        "text=Partners",
        "xpath=//a[@class='ui-tabs-anchor' and contains(@href,'_ovviewset.do_0008')]",
    ]
    clicked = False
    for sel in sels:
        try:
            loc = frame.locator(sel).first
            if loc.count():
                log("[nav] Clicking Partners tab")
                loc.click()
                clicked = True
                break
        except Exception:
            pass
    if not clicked:
        try:
            loc, ctx, _ = wait_find_in_any_frame(page, sels, timeout_ms=15000)
            loc.click()
            clicked = True
        except Exception:
            pass
    if not clicked:
        return False
    grid_sigs = [
        "xpath=//td[starts-with(@id,'GUIDE-PartnersTable-')]",
        "xpath=//td[@aria-label='Partner Function']",
        "xpath=//th[normalize-space(.)='Partner Function']",
    ]
    try:
        wait_find_in_any_frame(page, grid_sigs, timeout_ms=20000)
    except Exception:
        pass
    return True
def get_initial_reporter_name(frame):
    tr = _row_by_pf_in_partners(frame, ["Initial Reporter", "Initial Contact", "Initial Reporter/Contact"])
    if not tr:
        return ""
    return _cell_text_in_same_row(tr, "Name")
def get_facility_name_and_address(frame):
    tr = _row_by_pf_in_partners(frame, ["Facility", "Health Care Facility", "Healthcare Facility", "Plant"])
    if not tr:
        return ""
    name = _cell_text_in_same_row(tr, "Name")
    addr = _cell_text_in_same_row(tr, "Address") or _cell_text_in_same_row(tr, "address_short")
    if addr:
        addr = addr.replace(" / ", "\n")
    return f"{name}\n{addr}".strip()
def _partners_table(frame):
    t = frame.locator("xpath=//table[.//td[starts-with(@id,'GUIDE-PartnersTable-')]]").first
    return t if t.count() else None
def _row_by_pf_in_partners(frame, pf_names):
    log(f"[Partners] === Starting _row_by_pf_in_partners ===")
    log(f"[Partners] Looking for Partner Functions: {pf_names}")
    tbl = _partners_table(frame)
    if not tbl:
        log("[Partners] ERROR: Partners table not found")
        return None
    log("[Partners] Found Partners table")
    if isinstance(pf_names, str):
        pf_names = [pf_names]
    log("[Partners] EDIT MODE: Getting all rows with PartnerFunction cells")
    rows = tbl.locator(
        "xpath=.//tr[td[starts-with(@id,'GUIDE-PartnersTable-') and contains(@id,'-PartnerFunction')]]"
    )
    row_count = rows.count()
    log(f"[Partners] EDIT MODE: Found {row_count} rows with PartnerFunction cells")
    for i in range(row_count):
        log(f"[Partners] EDIT MODE: Examining row {i+1}/{row_count}")
        row = rows.nth(i)
        pf_value = _cell_text_in_same_row(row, "PartnerFunction")
        if not pf_value:
            log(f"[Partners] EDIT MODE: Row {i+1} - No PartnerFunction value found (skipping)")
            continue
        log(f"[Partners] EDIT MODE: Row {i+1} - Partner Function value: {pf_value!r}")
        pf_lower = pf_value.strip().lower()
        matched = False
        for want in pf_names:
            want_lower = want.strip().lower()
            if pf_lower == want_lower:
                log(f"[Partners] EDIT MODE: Row {i+1} - ✓ EXACT MATCH with {want!r}")
                matched = True
                return row
            if want_lower in pf_lower:
                log(f"[Partners] EDIT MODE: Row {i+1} - ✓ PARTIAL MATCH with {want!r} (in {pf_value!r})")
                matched = True
                return row
        if not matched:
            log(f"[Partners] EDIT MODE: Row {i+1} - No match")
    log("[Partners] ERROR: No matching row found")
    return None
def find_partners_frame(page, timeout_ms=10000, poll_ms=200):
    import time
    deadline = time.time() + (timeout_ms / 1000.0)
    sel = ("xpath=//td[starts-with(@id,'GUIDE-PartnersTable-') or "
           "starts-with(@id,'C21_W80_V81_btpartner_table')]")
    while time.time() < deadline:
        for fr in page.frames:
            try:
                loc = fr.locator(sel).first
                if loc.count():
                    return fr
            except Exception:
                pass
        time.sleep(poll_ms / 1000.0)
    return None
def debug_frames_for_partners(page):
    print("[Partners] Scanning frames for partners grid…")
    for i, fr in enumerate(page.frames):
        try:
            has = fr.locator("xpath=//td[starts-with(@id,'GUIDE-PartnersTable-') or starts-with(@id,'C21_W80_V81_btpartner_table')]").count()
            print(f"  [{i}] name={fr.name!r} url={fr.url!r}  matches={has}")
        except Exception as e:
            print(f"  [{i}] error: {e}")
def _cell_text_in_same_row(tr_loc, col_name):
    td = tr_loc.locator(
        f"xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-') and contains(@id,'-{col_name}')]"
    ).first
    if td.count():
        txt = clean(td.inner_text())
        if txt:
            return txt
        input_el = td.locator("xpath=.//input | .//select | .//textarea").first
        if input_el.count():
            try:
                val = input_el.input_value() if input_el.evaluate("e => e.tagName").lower() == 'input' else input_el.inner_text()
                if val:
                    return clean(val)
            except Exception:
                pass
        for attr in ["title", "aria-label", "value", "data-value"]:
            val = td.get_attribute(attr)
            if val:
                txt = clean(val)
                if txt:
                    return txt
    if col_name == "PartnerFunction":
        for variant in ["PartnerFunction", "Partner_Function", "PartnerFct", "Fct", "Function", "PartnerFunctionCode"]:
            td = tr_loc.locator(
                f"xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-') and contains(@id,'-{variant}')]"
            ).first
            if td.count():
                txt = clean(td.inner_text())
                if txt:
                    return txt
                input_el = td.locator("xpath=.//input | .//select | .//textarea").first
                if input_el.count():
                    try:
                        tag = input_el.evaluate("e => e.tagName").lower()
                        if tag == 'select':
                            selected = input_el.locator("xpath=.//option[@selected] | .//option[1]").first
                            if selected.count():
                                return clean(selected.inner_text())
                        else:
                            val = input_el.input_value()
                            if val:
                                return clean(val)
                    except Exception:
                        pass
                for attr in ["title", "aria-label", "value", "data-value"]:
                    val = td.get_attribute(attr)
                    if val:
                        txt = clean(val)
                        if txt:
                            return txt
    elif col_name == "Name":
        for variant in ["Name", "PartnerName", "Partner_Name"]:
            td = tr_loc.locator(
                f"xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-') and contains(@id,'-{variant}')]"
            ).first
            if td.count():
                txt = clean(td.inner_text())
                if txt:
                    return txt
                input_el = td.locator("xpath=.//input | .//textarea").first
                if input_el.count():
                    try:
                        val = input_el.input_value()
                        if val:
                            return clean(val)
                    except Exception:
                        pass
    elif col_name == "Address" or col_name == "address_short":
        for variant in ["Address", "address_short", "Address_Short", "PartnerAddress"]:
            td = tr_loc.locator(
                f"xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-') and contains(@id,'-{variant}')]"
            ).first
            if td.count():
                txt = clean(td.inner_text())
                if txt:
                    return txt
                textarea = td.locator("xpath=.//textarea | .//input").first
                if textarea.count():
                    try:
                        val = textarea.input_value()
                        if val:
                            return clean(val)
                    except Exception:
                        pass
    return ""
def _debug_list_pf_from_correct_table(frame):
    tbl = _partners_table(frame)
    if not tbl:
        print("[Partners] Could not find the Partners table (GUIDE-PartnersTable).")
        return
    cells = tbl.locator("xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-') and contains(@id,'-PartnerFunction')]")
    n = cells.count()
    print(f"[Partners] Partner table rows detected: {n}")
    for i in range(min(n, 30)):
        print("  -", clean(cells.nth(i).inner_text()))
def find_frame_with(page, selector, timeout_ms=10000, poll_ms=200):
    import time
    deadline = time.time() + (timeout_ms/1000.0)
    while time.time() < deadline:
        for fr in page.frames:
            try:
                if fr.locator(selector).first.count():
                    return fr
            except Exception:
                pass
        time.sleep(poll_ms/1000.0)
    return None
def click_tab_by_text(page, root_frame, text_or_href_fragment):
    sels = [
        f"xpath=//a[contains(@class,'ui-tabs-anchor') and normalize-space(.)='{text_or_href_fragment}']",
        f"text={text_or_href_fragment}",
        f"xpath=//a[contains(@class,'ui-tabs-anchor') and contains(@href,'{text_or_href_fragment}')]",
    ]
    for fr in [root_frame] + list(page.frames):
        for sel in sels:
            try:
                loc = fr.locator(sel).first
                if loc.count():
                    log(f"[nav] Clicking tab: {text_or_href_fragment}")
                    loc.click()
                    return fr
            except Exception:
                pass
    return None
_COMPLETE_RX = re.compile(
    r"\b(complete(?:d)?|closed\s*[-–]?\s*complete|closure\s*[-–]?\s*complete|final(?:ized)?|fully\s*resolved)\b",
    re.I,
)
_ID_PATTERNS = [
    r"\b\d{5,}\b",
    r"\bWI[-_ ]?\d{4,}\b",
    r"\bINV[-_ ]?\d{4,}\b",
    r"\bPA[-_ ]?\d{4,}\b",
    r"\bTXN[-_ ]?\d{4,}\b",
    r"\bAN[-_ ]?\d{4,}\b",
    r"\b[A-Z]{2,5}[-_ ]?\d{4,}\b",
]
def _find_first_match(patterns, s):
    for p in patterns:
        m = re.search(p, s or "", re.I)
        if m:
            return m.group(0)
    return None
def _row_combined_text(row):
    try:
        return clean(row.inner_text())
    except Exception:
        try:
            return (row.evaluate("n => n.textContent || ''") or "").strip()
        except Exception:
            return ""
def _row_txid(row):
    combined = _row_combined_text(row)
    txid = _find_first_match(_ID_PATTERNS, combined)
    if txid:
        return txid
    links = row.locator("xpath=.//a[@href]")
    for k in range(links.count()):
        href = (links.nth(k).get_attribute("href") or "")
        txid = _find_first_match(_ID_PATTERNS, href)
        if txid:
            return txid
        m = re.search(r"(?:id|no|number|case|txn|wi)[=:/#](\w[-\w]*)", href, re.I)
        if m:
            return m.group(1)
    cells = row.locator("xpath=.//th|.//td")
    for j in range(cells.count()):
        el = cells.nth(j)
        for attr in ("data-id", "data-transactionid", "data-txid", "data-key"):
            val = (el.get_attribute(attr) or "")
            txid = _find_first_match(_ID_PATTERNS, val)
            if txid:
                return txid
    return None
def _row_status_text(row):
    status_like = row.locator(
        "xpath=.//*[(self::td or self::th) and "
        " (contains(translate(@aria-label,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'status') or "
        "  contains(translate(@aria-label,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'state') or "
        "  contains(translate(@aria-label,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'result') or "
        "  contains(translate(@aria-label,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'outcome') or "
        "  contains(translate(@aria-label,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'disposition') or "
        "  contains(translate(@aria-label,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'resolution'))]"
    ).first
    if status_like.count():
        t = clean(status_like.inner_text()) or ""
        if t:
            return t
    hinted = row.locator("xpath=.//*[@title or @aria-label or @alt]")
    for k in range(hinted.count()):
        node = hinted.nth(k)
        for attr in ("title", "aria-label", "alt"):
            val = node.get_attribute(attr) or ""
            if _COMPLETE_RX.search(val):
                return val
    return _row_combined_text(row)
def _row_is_complete(row):
    t = (_row_status_text(row) or "").lower()
    if _COMPLETE_RX.search(t):
        return True
    return ("closed" in t and "complete" in t)
def _pli_table(frame):
    return frame.locator("xpath=//table[.//td[starts-with(@id,'GUIDE-ProductLineItemsTable-')]]").first
def _get_attr_or_text(node):
    try:
        text = clean(node.inner_text() or "")
    except Exception:
        text = ""
    if text:
        return text
    return clean(
        (node.get_attribute("title") or node.get_attribute("aria-label") or "")
    )
DEFAULT_PA_TEXT = (
    "Information provided to Medtronic indicated that the complaint device "
    "was not available for evaluation."
)
DEFAULT_INV_TEXT = (
    "Medtronic conducted an investigation based upon all received information. "
    "Without a product returned for evaluation a likely cause for the reported "
    "condition could not be established. Medtronic's assessment determined that "
    "manufacturing action is not required at this time. To ensure product oversight, "
    "this complaint report is incorporated into Medtronic's complaint monitoring and "
    "tracking system. The manufacturing records for each device are thoroughly "
    "reviewed to ensure the product meets its quality specifications. Should new "
    "information become available the file will be re-opened and the investigation "
    "summary will be amended as appropriate."
)
INV_ASSESSMENT_TAG = (
    "Medtronic's assessment determined that manufacturing action is not required at this time. "
    "To ensure product oversight, this complaint report is incorporated into Medtronic's complaint monitoring and tracking system. "
    "Should new information become available the file will be re-opened and the investigation summary will be amended as appropriate."
)
def _extract_investigation_body(text: str) -> str:
    if not text:
        return ""   
    lower = text.lower()
    start_phrase = "based on the evidence"
    idx = lower.find(start_phrase)
    if idx != -1:
        body_text = text[idx:]
    else:
        boiler_rx = re.compile(
            r'^\s*medtronic\s+conducted\s+an\s+investigation\s+based\s+upon\s+all\s+information\s+received\.?\s*',
            re.IGNORECASE
        )
        body_text = boiler_rx.sub("", text, count=1)
    lower_body = body_text.lower()
    dhr_phrase = "device history record"
    dhr_rel = lower_body.find(dhr_phrase)
    if dhr_rel == -1:
        return body_text.strip()
    snippet = body_text[:dhr_rel]
    last_dot = snippet.rfind(".")
    last_exc = snippet.rfind("!")
    last_q   = snippet.rfind("?")
    stop_rel = max(last_dot, last_exc, last_q)
    if stop_rel == -1:
        return body_text[:dhr_rel].strip()
    return body_text[:stop_rel + 1].strip()
def _strip_leading_based_on_evidence(text: str) -> str:
    if not text:
        return ""
    pattern = re.compile(
        r'^\s*based on the evidence available[,:]?\s*',
        re.IGNORECASE
    )
    m = pattern.match(text)
    if not m:
        return text.strip()
    rest = text[m.end():].lstrip()
    if not rest:
        return ""
    return rest[0].upper() + rest[1:]
def _postprocess_investigation_text(text: str) -> str:
    if not text:
        return ""
    body = _extract_investigation_body(text)
    body = _strip_leading_based_on_evidence(body)
    body = _replace_dashes_with_bullets(body)
    if INV_ASSESSMENT_TAG.lower() in body.lower():
        return body.strip()
    if body.endswith((".", "!", "?")):
        sep = " "
    else:
        sep = ". "
    return (body + sep + INV_ASSESSMENT_TAG).strip()
def _number_word(n: int) -> str:
    words = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five",
             6: "Six", 7: "Seven", 8: "Eight", 9: "Nine", 10: "Ten"}
    return words.get(n, str(n))
def _replace_dashes_with_bullets(text: str) -> str:
    if not text:
        return text
    return re.sub(r'^- ', '\u2022 ', text, flags=re.MULTILINE)
def _format_analysis_block(product_desc: str, summary: str, product_count: int = 1, include_lead: bool = True, has_image: bool = False) -> str:
    s = "" if summary is None else _normalize_text_preserve(summary)
    s = _replace_dashes_with_bullets(s)
    if not s.strip() or s.strip() == DEFAULT_PA_TEXT:
        return DEFAULT_PA_TEXT
    if not include_lead:
        return s.lstrip("\n")
    count_word = _number_word(product_count)
    desc = (product_desc or '').strip()
    if has_image:
        if product_count == 1:
            lead = (f"{count_word} {desc} and one picture were received for evaluation. "
                    f"Examination of the sample and picture is provided below.")
        else:
            lead = (f"{count_word} {desc} and one picture were received for evaluation. "
                    f"Examination of the samples and picture is provided below.")
    else:
        if product_count == 1:
            lead = f"{count_word} {desc} was received for evaluation. Examination of the sample is described below."
        else:
            lead = f"{count_word} {desc} were received for evaluation. Examination of the samples is described below."
    return lead + "\n\n" + s.lstrip("\n")
def get_current_activity_product_code(page) -> str:
    for fr in page.frames:
        try:
            td = fr.locator("xpath=//td[@id='bcTitle']").first
            if not td.count():
                continue
            title = (td.get_attribute("title") or td.inner_text() or "").strip()
            if not title:
                continue
            m = re.search(r'Product:([^,]+)', title)
            if m:
                code = m.group(1).strip()
                log(f"[bcTitle] Found product in header: {code!r} (title={title!r})")
                return code
        except Exception as e:
            log(f"[bcTitle] error reading bcTitle in frame {getattr(fr, 'name', '')}: {e}")
            continue
    log("[bcTitle] No product code found in any frame")
    return ""
def read_analysis_summary_and_product_for_txid(page, txid: str):
    ok = search_activities_for_id(page, txid)
    if not ok:
        log(f"[PA-SUMMARY] search_activities_for_id failed for txid={txid}")
        return "", "", ""
    prod_code = get_current_activity_product_code(page)
    log(f"[PA-SUMMARY] txid={txid} → bcTitle product={prod_code!r}")
    click_tab_by_text(page, page.main_frame, "Text Info") or \
    click_tab_by_text(page, page.main_frame, "_ovviewset.do_0006")
    raw_txt = (read_analysis_raw_text_for_current_pli(page) or "").strip()
    txt = (read_analysis_summary_for_current_pli(page) or "").strip()
    if not txt and raw_txt:
        log(f"[PA-SUMMARY] txid={txid} using raw_txt fallback for summary, len={len(raw_txt)}")
        txt = raw_txt
    if not txt:
        page.wait_for_timeout(500)
        txt2 = (read_analysis_summary_for_current_pli(page) or "").strip()
        if txt2:
            txt = txt2
        elif raw_txt:
            txt = raw_txt
    return txt, (prod_code or "").strip(), raw_txt
def read_investigation_summary_and_product_for_txid(page, txid: str):
    ok = search_activities_for_id(page, txid)
    if not ok:
        log(f"[INV-SUMMARY] search_activities_for_id failed for txid={txid}")
        return "", ""
    prod_code = get_current_activity_product_code(page)
    log(f"[INV-SUMMARY] txid={txid} → bcTitle product={prod_code!r}")
    click_tab_by_text(page, page.main_frame, "Text Info") or \
    click_tab_by_text(page, page.main_frame, "_ovviewset.do_0006")
    txt = read_investigation_summary_for_current_pli(page)
    if not txt:
        page.wait_for_timeout(500)
        txt = read_investigation_summary_for_current_pli(page)
    return (txt or "").strip(), (prod_code or "").strip()
def _normalize_text_preserve(s: str) -> str:
    if s is None:
        return ""
    s = s.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    lines = s.split("\n")
    lines = [re.sub(r"[ \t]+$", "", ln) for ln in lines]
    s = "\n".join(lines)
    return _strip_boilerplate_sentences(s)
def _safe_td_text(td, preserve=False):
    if not td or not td.count():
        return ""
    try:
        if preserve:
            raw = td.evaluate("n => n.textContent || ''") or ""
            return _normalize_text_preserve(raw)
        raw = td.evaluate("""
            (node) => {
                const a = node.querySelector("a[id*='text_table'][id*='lines']");
                if (a) {
                    const t = a.getAttribute('title') || a.getAttribute('aria-label');
                    if (t && t.trim()) return t;
                }
                return node.textContent || '';
            }
        """) or ""
    except Exception:
        try:
            raw = td.evaluate("n => n.textContent || ''") or ""
        except Exception:
            raw = ""
    return _normalize_text(raw) if not preserve else _normalize_text_preserve(raw)
def read_all_products(page, root_frame):
    click_tab_by_text(page, root_frame, "Product Line Items") or \
    click_tab_by_text(page, root_frame, "_ovviewset.do_0002")
    fr = find_frame_with(page, "xpath=//td[starts-with(@id,'GUIDE-ProductLineItemsTable-')]")
    if fr:
        tbl = _pli_table(fr)
        if tbl and tbl.count():
            rows = tbl.locator(
                "xpath=.//tr[td[starts-with(@id,'GUIDE-ProductLineItemsTable-') and contains(@id,'-Product')]]"
            )
            n = rows.count()
            out = []
            for i in range(n):
                row = rows.nth(i)
                complaint_val = ""
                complaint_td = row.locator(
                    "xpath=.//td[starts-with(@id,'GUIDE-ProductLineItemsTable-') "
                    "and contains(@id,'-Complaint')]"
                ).first
                if complaint_td.count():
                    span = complaint_td.locator("xpath=.//span").first
                    if span.count():
                        complaint_val = clean(
                            span.get_attribute("title")
                            or span.inner_text()
                        )
                    else:
                        complaint_val = clean(complaint_td.inner_text())
                    if not complaint_val:
                        input_el = complaint_td.locator("xpath=.//input | .//select").first
                        if input_el.count():
                            try:
                                tag = input_el.evaluate("e => e.tagName").lower()
                                if tag == 'select':
                                    selected = input_el.locator("xpath=.//option[@selected] | .//option[1]").first
                                    if selected.count():
                                        complaint_val = clean(selected.inner_text())
                                else:
                                    val = input_el.input_value()
                                    if val:
                                        complaint_val = clean(val)
                                log(f"[PLI] Row {i+1}/{n}: Got Complaint value from input element: {complaint_val!r}")
                            except Exception as e:
                                log(f"[PLI] Row {i+1}/{n}: Error reading Complaint input: {e}")
                complaint_norm = (complaint_val or "").strip().lower()
                log(f"[PLI] Row {i+1}/{n}: Complaint={complaint_val!r} (normalized={complaint_norm!r})")
                if complaint_norm.startswith("no"):
                    log(f"[PLI] Skipping product row {i+1}/{n} because Complaint={complaint_val!r}")
                    continue
                prod_cell = row.locator(
                    f"xpath=.//td[starts-with(@id,'GUIDE-ProductLineItemsTable-') and {ENDS_WITH_PRODUCT}]"
                ).first
                desc_cell = row.locator(
                    f"xpath=.//td[starts-with(@id,'GUIDE-ProductLineItemsTable-') and {ENDS_WITH_DESCRIPTION}]"
                ).first
                pid = ""
                if prod_cell.count():
                    a = prod_cell.locator("xpath=.//a[contains(@id,'ordered_prod')]").first
                    if a.count():
                        pid = _get_attr_or_text(a)
                    if not pid:
                        pid = _get_attr_or_text(prod_cell)
                if pid and pid.strip().lower() == "open input help":
                    log(f"[PLI] Ignoring 'Open Input Help' as product ID in row {i+1}")
                    pid = ""
                if pid and not re.search(r"[A-Za-z]", pid):
                    log(f"[PLI] Filtering out numeric-only product ID: {pid!r}")
                    pid = ""
                pdesc = ""
                if desc_cell.count():
                    a = desc_cell.locator("xpath=.//a").first
                    pdesc = _get_attr_or_text(a) if a.count() else _get_attr_or_text(desc_cell)
                pcode = pid or extract_product_code(pdesc)
                sn_val = ""
                sn_cell = row.locator(
                    "xpath=.//td[starts-with(@id,'GUIDE-ProductLineItemsTable-') and "
                    "(contains(@id,'-SN') or contains(@id,'-SerialNumber') or contains(@id,'-Serial')) "
                    "and not(contains(@id,'SNValid'))]"
                ).first
                if sn_cell.count():
                    sn_val = clean(sn_cell.inner_text())
                    if not sn_val:
                        sn_val = clean(sn_cell.get_attribute("title") or sn_cell.get_attribute("aria-label") or "")
                    if not sn_val:
                        input_el = sn_cell.locator("xpath=.//input").first
                        if input_el.count():
                            try:
                                val = input_el.input_value()
                                if val:
                                    sn_val = clean(val)
                            except Exception:
                                pass
                lot_val = ""
                lot_cell = row.locator(
                    "xpath=.//td[starts-with(@id,'GUIDE-ProductLineItemsTable-') and "
                    "(contains(@id,'-Lot') or contains(@id,'-LotNumber')) "
                    "and not(contains(@id,'LotValid'))]"
                ).first
                if lot_cell.count():
                    lot_val = clean(lot_cell.inner_text())
                    if not lot_val:
                        lot_val = clean(lot_cell.get_attribute("title") or lot_cell.get_attribute("aria-label") or "")
                    if not lot_val:
                        input_el = lot_cell.locator("xpath=.//input").first
                        if input_el.count():
                            try:
                                val = input_el.input_value()
                                if val:
                                    lot_val = clean(val)
                            except Exception:
                                pass
                if pid or pdesc or sn_val or lot_val:
                    out.append({
                        "id": pid,
                        "desc": pdesc,
                        "code": pcode,
                        "sn": sn_val,
                        "lot": lot_val,
                        "complaint": complaint_val or "",
                    })
            if out:
                return out
    fr = find_frame_with(page, "xpath=//*[contains(@id,'btadmini_table')]")
    if not fr:
        return []
    rows = fr.locator("xpath=.//tr[.//a[contains(@id,'ordered_prod')]]")
    n = rows.count()
    out = []
    for i in range(n):
        row = rows.nth(i)
        complaint_val = ""
        complaint_span = row.locator(
            "xpath=.//span[contains(@id,'zcomplaint') or @aria-label='Complaint']"
        ).first
        if complaint_span.count():
            complaint_val = clean(
                complaint_span.get_attribute("title") or complaint_span.inner_text()
            )
        if not complaint_val:
            complaint_input = row.locator(
                "xpath=.//input[contains(@id,'zcomplaint')] | .//select[contains(@id,'zcomplaint')]"
            ).first
            if complaint_input.count():
                try:
                    tag = complaint_input.evaluate("e => e.tagName").lower()
                    if tag == 'select':
                        selected = complaint_input.locator("xpath=.//option[@selected] | .//option[1]").first
                        if selected.count():
                            complaint_val = clean(selected.inner_text())
                    else:
                        val = complaint_input.input_value()
                        if val:
                            complaint_val = clean(val)
                    log(f"[PLI-btadmini] Row {i+1}/{n}: Got Complaint from input: {complaint_val!r}")
                except Exception as e:
                    log(f"[PLI-btadmini] Row {i+1}/{n}: Error reading Complaint input: {e}")
        complaint_norm = (complaint_val or "").strip().lower()
        log(f"[PLI-btadmini] Row {i+1}/{n}: Complaint={complaint_val!r} (normalized={complaint_norm!r})")
        if complaint_norm.startswith("no"):
            log(f"[PLI-btadmini] Skipping product row {i+1}/{n} because Complaint={complaint_val!r}")
            continue
        ordered_link = row.locator("xpath=.//a[contains(@id,'ordered_prod')]").first
        pid = _get_attr_or_text(ordered_link) if ordered_link.count() else ""
        if pid and pid.strip().lower() == "open input help":
            log(f"[PLI-btadmini] Ignoring 'Open Input Help' as product ID in row {i+1}")
            pid = ""
        pdesc = ""
        try:
            if ordered_link.count():
                maybe_desc = ordered_link.locator("xpath=ancestor::td[1]/following-sibling::td[1]").first
                if maybe_desc.count():
                    pdesc = clean(maybe_desc.inner_text())
        except Exception:
            pass
        pcode = pid or extract_product_code(pdesc)
        sn_val = ""
        lot_val = ""
        if pid or pdesc:
            out.append({
                "id": pid,
                "desc": pdesc,
                "code": pcode,
                "sn": sn_val,
                "lot": lot_val,
                "complaint": complaint_val or "",
            })
    return out
def _dates_table(frame):
    return frame.locator("xpath=//table[.//td[starts-with(@id,'GUIDE-DatesTable')]]").first
def get_event_date(page):
    click_tab_by_text(page, page.main_frame, "Dates") or \
    click_tab_by_text(page, page.main_frame, "_ovviewset.do_0003")
    fr = find_frame_with(page, "xpath=//td[starts-with(@id,'GUIDE-DatesTable')]")
    if not fr:
        log("[EventDate] Could not find frame with Dates table")
        return ""
    tbl = _dates_table(fr)
    if not tbl or not tbl.count():
        log("[EventDate] Dates table not found in frame")
        return ""
    date_type_variations = [
        "Event Date",
        "Event date",
        "event date",
        "Date of Event",
        "Incident Date",
    ]
    row = None
    for date_type in date_type_variations:
        row = tbl.locator(
            f"xpath=.//tr[td[starts-with(@id,'GUIDE-DatesTable') and contains(@id,'-DateType') and normalize-space(.)='{date_type}']]"
        ).first
        if row.count():
            log(f"[EventDate] Found row with DateType={date_type!r}")
            break    
    if not row or not row.count():
        log("[EventDate] No row found with Event Date type")
        row = tbl.locator(
            "xpath=.//tr[td[starts-with(@id,'GUIDE-DatesTable') and contains(@id,'-DateType') "
            "and contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'), 'event')]]"
        ).first
        if not row.count():
            return ""
    cell = row.locator(
        "xpath=.//td[starts-with(@id,'GUIDE-DatesTable') and contains(@id,'-DateFrom')]"
    ).first
    if not cell.count():
        log("[EventDate] DateFrom cell not found in row")
        return ""
    event_date = clean(cell.inner_text())
    if not event_date:
        input_el = cell.locator("xpath=.//input | .//textarea").first
        if input_el.count():
            try:
                val = input_el.input_value()
                if val:
                    event_date = clean(val)
            except Exception:
                pass
    if not event_date:
        for attr in ["title", "aria-label", "value", "data-value"]:
            val = cell.get_attribute(attr)
            if val:
                event_date = clean(val)
                if event_date:
                    break
    log(f"[EventDate] Extracted event date: {event_date!r}")
    return event_date
def _aer_table(frame):
    return frame.locator("xpath=//table[.//td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-')]]").first
def _aer_row_by_type(tbl, *type_fragments):
    for t in type_fragments:
        row = tbl.locator(
            f"xpath=.//tr[td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-') and contains(@id,'-ExtReferenceType') and normalize-space(.)='{t}']]"
        ).first
        if row.count():
            return row
    for t in type_fragments:
        low = t.lower()
        row = tbl.locator(
            "xpath=.//tr[td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-') and contains(@id,'-ExtReferenceType') "
            f"and contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'), '{low}')]]"
        ).first
        if row.count():
            return row
    return None
def _aer_rows_by_type(tbl, *type_fragments):
    rows = []
    for t in type_fragments:
        cand = tbl.locator(
            "xpath=.//tr[td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-') "
            "and contains(@id,'-ExtReferenceType') and normalize-space(.)=$t]]"
        ).filter(has_text=t)
        for i in range(cand.count()):
            rows.append(cand.nth(i))
    if rows:
        return rows
    for t in type_fragments:
        low = t.lower()
        cand = tbl.locator(
            "xpath=.//tr[td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-') "
            "and contains(@id,'-ExtReferenceType') and "
            "contains(translate(normalize-space(.),"
            "'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),"
            f" '{low}')]]"
        )
        for i in range(cand.count()):
            rows.append(cand.nth(i))
    return rows
def _aer_comment_from_row(row):
    try:
        td = row.locator(
            "xpath=.//td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-') "
            "and contains(@id,'-ExtReferenceText')]"
        ).first
        if td.count():
            return clean(td.inner_text())
    except Exception:
        pass
    return ""
def _aer_number_from_row(row):
    for suffix in ("-ExtReferenceNumber", "-ExtReferenceID", "reference_number", "reference_id"):
        td = row.locator(
            f"xpath=.//td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-') and contains(@id,'{suffix}')]"
        ).first
        if td.count():
            return clean(td.inner_text())
    return ""
def read_external_refs(page, root_frame):
    owner = click_tab_by_text(page, root_frame, "Additional External References") or \
            click_tab_by_text(page, root_frame, "_ovviewset.do_0013")
    fr = find_frame_with(page, "xpath=//td[starts-with(@id,'GUIDE-AdditionalExternalReferencesTable-')]")
    if not fr:
        return {"rb_reference": "", "report_number": ""}
    tbl = _aer_table(fr)
    if not tbl or not tbl.count():
        return {"rb_reference": "", "report_number": ""}
    out = {"rb_reference": "", "report_number": ""}
    rb_row = _aer_row_by_type(tbl, "RB Reference", "RB", "RBReference")
    if rb_row:
        out["rb_reference"] = _aer_number_from_row(rb_row)
    rep_row = _aer_row_by_type(tbl, "MPXR")
    if rep_row:
        out["report_number"] = _aer_number_from_row(rep_row)
    ex_row = _aer_row_by_type(
        tbl,
        "SAP ECC Service & Repair",
        "SAP ECC Service and Repair",   # optional variant
        "SAP ECC",                      # optional fallback
        "ECC Service & Repair"          # optional fallback
    )
    if ex_row:
        out["ex_ref"] = _aer_number_from_row(ex_row)
    contact_rows = _aer_rows_by_type(
        tbl,
        "External Contact", "ExternalContact", "Ext Contact", "Ext. Contact"
    )
    contacts = []
    for row in contact_rows or []:
        num = _aer_number_from_row(row)
        txt = _aer_comment_from_row(row)
        contacts.append({
            "number": num,
            "text": txt,
        })
    out["external_contacts"] = contacts
    return out
def _find_latest_text_table(page):
    tables = []
    for fr in page.frames:
        try:
            scope = fr.locator("xpath=//div[contains(@class,'th-clr-cnt-bottom')]")
            scope = scope if scope.count() else fr
            tds = scope.locator("xpath=.//table[.//td[starts-with(@id,'GUIDE-TextInfoTable-')]]")
            n = tds.count()
            for i in range(n):
                tables.append((fr, tds.nth(i)))
        except Exception:
            pass
    return tables[-1] if tables else (None, None)
def read_event_description(page, root_frame):
    click_tab_by_text(page, root_frame, "Text Info") or \
    click_tab_by_text(page, root_frame, "_ovviewset.do_0006")
    frames = active_content_frames(page, {"content_frame_name_regex": "WorkAreaFrame1"})
    tbl = None
    chosen_frame = None
    for fr in frames:
        try:
            fr.wait_for_selector("xpath=//td[starts-with(@id,'GUIDE-TextInfoTable-')]", timeout=3000)
            candidate = _find_latest_text_table_in(fr)
            if candidate:
                tbl, chosen_frame = candidate, fr
                break
        except Exception:
            continue
    if not (tbl and chosen_frame):
        fr, tbl_fallback = _find_latest_text_table(page)
        if not (fr and tbl_fallback and tbl_fallback.count()):
            return ""
        chosen_frame, tbl = fr, tbl_fallback
    want_types = [
        "Incident description", "Incident description / Reason for report",
        "Reason for report", "Description of Event", "Event Description",
        "Narrative", "HCP Narrative", "Event narrative", "Incident Narrative"
    ]
    row = None
    for t in want_types:
        cand = tbl.locator(
            "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-TextType') and normalize-space(.)=$t]]"
        ).filter(has_text=t).first
        if cand.count():
            row = cand
            break
    if not (row and row.count()):
        for t in want_types:
            low = t.lower()
            cand = tbl.locator(
                "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-TextType') "
                f"and contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'), '{low}')]]"
            ).first
            if cand.count():
                row = cand
                break
    if not (row and row.count()):
        row = tbl.locator("xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-Text')]]").first
        if not row.count():
            log("[Text] No row with -Text found")
            return ""
    text_td = row.locator(
        "xpath=.//td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-Text') and not(contains(@id,'-TextType'))]"
    ).first
    if not text_td.count():
        log("[Text] No -Text cell present in row")
        return ""
    _suppress_clicks_enable(chosen_frame)
    try:
        desc = _safe_td_text(text_td)
    finally:
        _suppress_clicks_disable(chosen_frame)
    log(f"[Text] Event description read from TD (no click), length={len(desc)}")
    return desc
def _textinfo_signature(page):
    fr, tbl = _find_latest_analysis_table_nearby(page)
    if not (fr and tbl and tbl.count()):
        return ""
    try:
        s = tbl.evaluate("t => (t.innerText || '').slice(0, 800)")
    except Exception:
        s = tbl.inner_text()[:800] if tbl.count() else ""
    return s
def wait_for_textinfo_change(page, previous_sig, timeout=10000):
    start = time.time()
    while time.time() - start < (timeout/1000.0):
        sig = _textinfo_signature(page)
        if sig and sig != previous_sig:
            return True
        time.sleep(0.15)
    return False
def _find_leftnav_frame(page):
    for fr in page.frames:
        try:
            if fr.locator("xpath=//*[contains(@class,'left-nav')]").first.count():
                return fr
        except Exception:
            pass
    for fr in page.frames:
        try:
            if fr.locator("css=a[data-trans-id], a[data-transId], a[data-transid], a.GUIDE-sideNav, a.GUIDE-sideNavLink").first.count():
                return fr
        except Exception:
            pass
    return None
def _section_class(section_text: str) -> str:
    return {
        "Product Analysis": "ProductAnalysis",
        "Investigations":   "Investigation",
        "Product Line Items":"PLI",
    }.get(section_text, re.sub(r"\W+", "", section_text))
def _section_anchor_xpath(section: str) -> str:
    cls = _section_class(section)
    return (
        "xpath=//*[contains(@class,'left-nav')]"
        f"//div[contains(@class,'{cls}')]"
        "/following-sibling::*[contains(@class,'clicker')][1]"
        "/following-sibling::*[contains(@class,'data-wrapper')][1]"
        "//a[(contains(@class,'GUIDE-sideNav') or contains(@class,'GUIDE-sideNavLink') "
        "     or @data-trans-id or @data-transId or @data-transid)]"
    )
def _leftnav_section_header(fr, section_text: str):
    cls = _section_class(section_text)
    header = fr.locator(
        f"xpath=//*[contains(@class,'left-nav')]//div[contains(@class,'{cls}')]"
    ).first
    return header if header.count() else None
def _leftnav_section_container(fr, section_text: str):
    header = _leftnav_section_header(fr, section_text)
    if not header:
        return None
    clicker = header.locator("xpath=following-sibling::div[contains(@class,'clicker')][1]").first
    container = header.locator("xpath=following-sibling::div[contains(@class,'data-wrapper')][1]").first
    if container.count() == 0 and clicker.count():
        robust_click(clicker, fr)
        try:
            fr.wait_for_timeout(250)
        except Exception:
            pass
        container = header.locator("xpath=following-sibling::div[contains(@class,'data-wrapper')][1]").first
    return container if container.count() else None
def _anchors_in_section(fr, section_text: str):
    anchors = fr.locator(_section_anchor_xpath(section_text))
    if anchors.count():
        return anchors
    container = _leftnav_section_container(fr, section_text)
    return container.locator("xpath=.//a[(contains(@class,'GUIDE-sideNav') or contains(@class,'GUIDE-sideNavLink') "
                             "           or @data-trans-id or @data-transId or @data-transid)]") if container else fr.locator("xpath=//*[(0=1)]")
def _enumerate_section_items(fr, section_text: str):
    anchors = _anchors_in_section(fr, section_text)
    n = anchors.count()
    items = []
    for i in range(n):
        a = anchors.nth(i)
        try:
            txt = (a.inner_text().strip() or (a.get_attribute("title") or "").strip())
        except Exception:
            txt = (a.get_attribute("title") or "").strip()
        did = a.get_attribute("data-trans-id") or a.get_attribute("data-transId") or a.get_attribute("data-transid") or ""
        code = extract_product_code(txt)
        items.append({"i": i, "text": txt, "code": code, "data_id": (did or "").strip()})
    log(f"[LeftNav:{section_text}] anchors detected: {n}")
    return items
def _scan_pa_anywhere(page, section_text: str):
    found = []
    for fr in page.frames:
        try:
            anchors = fr.locator(
                "xpath=//*[contains(@class,'left-nav')]"
                "//div[contains(@class,'ProductAnalysis')]"
                "/following-sibling::*[contains(@class,'clicker')][1]"
                "/following-sibling::*[contains(@class,'data-wrapper')][1]"
                "//a[(contains(@class,'GUIDE-sideNav') or contains(@class,'GUIDE-sideNavLink') "
                "     or @data-trans-id or @data-transId or @data-transid)]"
            )
            n = anchors.count()
            if n == 0:
                anchors = fr.locator(
                    "xpath=//*[contains(@class,'left-nav')]"
                    "//a[(contains(@class,'GUIDE-sideNav') or contains(@class,'GUIDE-sideNavLink') "
                    "     or @data-trans-id or @data-transId or @data-transid)]"
                )
                n = anchors.count()
            for i in range(n):
                a = anchors.nth(i)
                try:
                    txt = (a.inner_text().strip() or (a.get_attribute('title') or '')).strip()
                except Exception:
                    txt = (a.get_attribute('title') or '').strip()
                did = (a.get_attribute('data-trans-id') or a.get_attribute('data-transId') or a.get_attribute('data-transid') or '').strip()
                found.append((fr, a, txt, did))
        except Exception:
            continue
    return found
def _enumerate_pa_items(fr):
    return _enumerate_section_items(fr, "Product Analysis")
def _content_frame(page):
    for fr in page.frames:
        if (fr.name or "") == "WorkAreaFrame1":
            return fr
    return page.main_frame
def _remove_product_desc_from_event(event_text: str, products: list) -> str:
    if not event_text or not products:
        return event_text
    text = event_text
    descriptions_to_remove = []
    for p in products:
        desc = (p.get("desc") or "").strip()
        if desc and len(desc) > 10:  # Only remove substantial descriptions
            descriptions_to_remove.append(desc)
    for desc in descriptions_to_remove:
        escaped_desc = re.escape(desc)
        pattern = r'\(\s*' + escaped_desc + r'\s*\)'
        text = re.sub(pattern, '', text, flags=re.IGNORECASE)
    text = re.sub(r'\s{2,}', ' ', text)
    text = re.sub(r'\s+([.,;!?])', r'\1', text)  # Remove space before punctuation
    return text.strip()
def _ensure_section_expanded(page, section: str):
    fr = _find_leftnav_frame(page)
    if not fr:
        return
    if section == "Product Analysis":
        _pa_try_expand(fr)
        try:
            fr.wait_for_selector(_section_anchor_xpath("Product Analysis"), timeout=1200)
        except Exception:
            pass
        return
    cls = _section_class(section)
    header = fr.locator(
        f"xpath=//*[contains(@class,'left-nav')]//*[contains(@class,'{cls}')]"
    ).first
    if not header.count():
        return
    clicker = header.locator("xpath=following-sibling::*[contains(@class,'clicker')][1]").first
    container = header.locator("xpath=following-sibling::*[contains(@class,'data-wrapper')][1]").first
    need_click = True
    if container.count():
        try:
            need_click = container.evaluate("n => getComputedStyle(n).display === 'none'")
        except Exception:
            pass
    if need_click and clicker.count():
        robust_click(clicker, fr)
        try: fr.wait_for_timeout(250)
        except Exception: pass
def list_side_nav_items(page, section: str):
    fr = _find_leftnav_frame(page)
    if not fr:
        return []
    _ensure_section_expanded(page, section)
    anchors = fr.locator(_section_anchor_xpath(section))
    n = anchors.count()
    out = []
    for i in range(n):
        el = anchors.nth(i)
        try:
            t = clean(el.inner_text())
            if not t:
                t = (el.get_attribute("title") or "").strip()
        except Exception:
            t = (el.get_attribute("title") or "").strip()
        code = extract_product_code(t)
        out.append({"text": t, "code": code, "el": el, "frame": fr})
    return out
def _find_latest_analysis_table_nearby(page):
    candidates = []
    for fr in page.frames:
        try:
            scope = fr.locator("xpath=//div[contains(@class,'th-clr-cnt-bottom')]")
            scope = scope if scope.count() else fr
            tds = scope.locator("xpath=.//table[.//td[starts-with(@id,'GUIDE-TextInfoTable-')]]")
            n = tds.count()
            for i in range(n):
                candidates.append((fr, tds.nth(i)))
        except Exception:
            pass
    return candidates[-1] if candidates else (None, None)
def _read_detail_textarea_from_frame(fr, preserve_format=False, timeout_ms=2500):
    sels = [
        "css=textarea[id$='_text_lines']",
        "css=textarea[id*='text_lines']",
        "xpath=//div[contains(@class,'th-ta-container')]//textarea",
        "xpath=//textarea[contains(@id,'_text_lines') or contains(@id,'text_lines')]",
    ]
    for sel in sels:
        try:
            ta = fr.locator(sel).first
            if ta.count():
                try:
                    ta.wait_for(state="visible", timeout=timeout_ms)
                except Exception:
                    pass
                try:
                    val = ta.input_value()  # best for textarea
                except Exception:
                    val = ta.evaluate("n => n.value || n.textContent || ''") or ""
                val = val.strip()
                if val:
                    return _normalize_text_preserve(val) if preserve_format else _normalize_text(val)
        except Exception:
            continue
    return ""
def xpath_literal(s: str) -> str:
    if s is None:
        return "''"
    if "'" not in s:
        return f"'{s}'"
    if '"' not in s:
        return f'"{s}"'
    parts = s.split("'")
    concat_parts = []
    for i, part in enumerate(parts):
        if part:
            concat_parts.append(f"'{part}'")
        if i < len(parts) - 1:
            concat_parts.append('"\'"')  # a literal single quote
    return "concat(" + ", ".join(concat_parts) + ")"
def _read_textarea_via_label_for(fr, preserve_format=False):
    lab = fr.locator("xpath=//label[contains(@id,'text_lines') and @for]").first
    if lab.count():
        tid = (lab.get_attribute("for") or "").strip()
        if tid:
            ta = fr.locator(f"xpath=//textarea[@id={xpath_literal(tid)}]").first
            if ta.count():
                try:
                    val = ta.input_value()
                except Exception:
                    val = ta.evaluate("n => n.value || n.textContent || ''") or ""
                val = (val or "").strip()
                if val:
                    return _normalize_text_preserve(val) if preserve_format else _normalize_text(val)
    return ""
def read_text_by_labels(page, wanted_labels, *, preserve_format=False):
    log(f"[TextInfo] === Starting read_text_by_labels ===")
    log(f"[TextInfo] Wanted labels: {wanted_labels}")
    log(f"[TextInfo] Preserve format: {preserve_format}")
    fr, tbl = _find_latest_analysis_table_nearby(page)
    if not (fr and tbl and tbl.count()):
        log("[TextInfo] ERROR: No TextInfo table found")
        return None
    log(f"[TextInfo] Found TextInfo table in frame: {getattr(fr, 'name', '')} url={getattr(fr, 'url', '')[:100]}")
    row = None
    log("[TextInfo] EDIT MODE: Starting row iteration to find matching TextType")
    all_rows = tbl.locator(
        "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-TextType')]]"
    )
    row_count = all_rows.count()
    log(f"[TextInfo] EDIT MODE: Found {row_count} rows with TextType cells")
    for i in range(row_count):
        log(f"[TextInfo] EDIT MODE: Examining row {i+1}/{row_count}")
        candidate_row = all_rows.nth(i)
        type_td = candidate_row.locator(
            "xpath=.//td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-TextType')]"
        ).first
        if not type_td.count():
            log(f"[TextInfo] EDIT MODE: Row {i+1} - No TextType cell found (skipping)")
            continue
        type_value = ""
        select = type_td.locator("xpath=.//select").first
        if select.count():
            log(f"[TextInfo] EDIT MODE: Row {i+1} - Found <select> element (EDIT MODE DETECTED)")
            try:
                selected = select.locator("xpath=.//option[@selected]").first
                if selected.count():
                    type_value = clean(selected.inner_text())
                    log(f"[TextInfo] EDIT MODE: Row {i+1} - Got value from @selected option: {type_value!r}")
                else:
                    type_value = clean(select.evaluate("""
                        el => {
                            const idx = el.selectedIndex;
                            return idx >= 0 && el.options[idx] ? el.options[idx].text : '';
                        }
                    """))
                    log(f"[TextInfo] EDIT MODE: Row {i+1} - Got value from selectedIndex: {type_value!r}")
            except Exception as e:
                log(f"[TextInfo] EDIT MODE: Row {i+1} - ERROR reading select: {e}")
        else:
            log(f"[TextInfo] EDIT MODE: Row {i+1} - No <select> element found (read-only mode)")
        if not type_value:
            type_value = clean(type_td.inner_text())
            log(f"[TextInfo] EDIT MODE: Row {i+1} - Got value from inner_text: {type_value!r}")
        if not type_value:
            log(f"[TextInfo] EDIT MODE: Row {i+1} - No TextType value found (skipping)")
            continue
        type_lower = type_value.strip().lower()
        log(f"[TextInfo] EDIT MODE: Row {i+1} - Comparing {type_lower!r} against wanted labels")
        matched = False
        for label in wanted_labels:
            label_lower = label.strip().lower()
            if type_lower == label_lower:
                log(f"[TextInfo] EDIT MODE: Row {i+1} - ✓ EXACT MATCH with {label!r}")
                row = candidate_row
                matched = True
                break
            elif label_lower in type_lower:
                log(f"[TextInfo] EDIT MODE: Row {i+1} - ✓ PARTIAL MATCH with {label!r}")
                row = candidate_row
                matched = True
                break
        if not matched:
            log(f"[TextInfo] EDIT MODE: Row {i+1} - No match")
        if row:
            break
    if row:
        log("[TextInfo] EDIT MODE: Successfully found row via iteration")
    else:
        log("[TextInfo] EDIT MODE: No row found via iteration, falling back to XPath text matching")
    if not row:
        log("[TextInfo] READ-ONLY MODE: Trying XPath exact text matching")
        for t in wanted_labels:
            cand = tbl.locator(
                "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') "
                "and contains(@id,'-TextType') and normalize-space(.)=" + xpath_literal(t) + "]]"
            ).first
            if cand.count():
                log(f"[TextInfo] READ-ONLY MODE: Found row with exact text match for {t!r}")
                row = cand
                break
    if not row:
        log("[TextInfo] READ-ONLY MODE: Trying XPath case-insensitive text matching")
        for t in wanted_labels:
            low = t.lower()
            cand = tbl.locator(
                "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') "
                "and contains(@id,'-TextType') and "
                f"contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'), '{low}')]]"
            ).first
            if cand.count():
                log(f"[TextInfo] READ-ONLY MODE: Found row with case-insensitive match for {t!r}")
                row = cand
                break
    if not row:
        log("[TextInfo] ERROR: No row found with any method")
        return None
    log("[TextInfo] Row found! Extracting text content...")
    td = row.locator(
        "xpath=.//td[starts-with(@id,'GUIDE-TextInfoTable-') "
        "and contains(@id,'-Text') and not(contains(@id,'-TextType'))]"
    ).first
    if not td.count():
        log("[TextInfo] ERROR: No Text cell found in row")
        return None
    log("[TextInfo] Found Text cell, attempting extraction methods...")
    a = td.locator("xpath=.//a[contains(@id,'text_table') and contains(@id,'lines')]").first
    if a.count():
        full = (a.get_attribute('title') or a.get_attribute('aria-label') or '').strip()
        if full:
            log(f"[TextInfo] ✓ SUCCESS: Got text from <a> aria-label/title, length={len(full)}")
            return _normalize_text_preserve(full) if preserve_format else _normalize_text(full)
        else:
            log("[TextInfo] Found <a> element but title/aria-label was empty")
    else:
        log("[TextInfo] No <a> element with text_table/lines found")
    txt = _safe_td_text(td, preserve=preserve_format)
    if txt:
        log(f"[TextInfo] ✓ SUCCESS: Got text from _safe_td_text, length={len(txt)}")
        return txt
    else:
        log("[TextInfo] _safe_td_text returned empty")
    try:
        log("[TextInfo] Trying click-to-load detail panel textarea fallback...")
        clicked = robust_click(row, fr) or robust_click(td, fr)
        if clicked:
            try:
                fr.locator("css=textarea[id*='text_lines']").first.wait_for(state="attached", timeout=3000)
            except Exception:
                pass
            fr.wait_for_timeout(150)
            detail = _read_detail_textarea_from_frame(fr, preserve_format=preserve_format, timeout_ms=3000)
            if not detail:
                detail = _read_textarea_via_label_for(fr, preserve_format=preserve_format)
            if detail:
                log(f"[TextInfo] ✓ SUCCESS: Got text from detail panel textarea, length={len(detail)}")
                return detail
            else:
                log("[TextInfo] Detail panel textarea not found / empty after click")
    except Exception as e:
        log(f"[TextInfo] Detail panel click-to-load fallback failed: {e}")
    input_el = td.locator("xpath=.//textarea | .//input").first
    if input_el.count():
        log("[TextInfo] Found textarea/input element (EDIT MODE)")
        try:
            val = input_el.input_value()
            if val:
                result = _normalize_text_preserve(val) if preserve_format else _normalize_text(val)
                log(f"[TextInfo] ✓ SUCCESS: Got text from input element, length={len(result)}")
                return result
            else:
                log("[TextInfo] input_value() returned empty")
        except Exception as e:
            log(f"[TextInfo] ERROR reading input element: {e}")
    else:
        log("[TextInfo] No textarea/input element found")
    wysiwyg = td.locator("xpath=.//div[contains(@class,'th-wysi') or contains(@class,'th-txt')]").first
    if wysiwyg.count():
        log("[TextInfo] Found WYSIWYG div (EDIT MODE)")
        try:
            raw = wysiwyg.evaluate("n => n.textContent || ''")
            if raw:
                result = _normalize_text_preserve(raw) if preserve_format else _normalize_text(raw)
                log(f"[TextInfo] ✓ SUCCESS: Got text from WYSIWYG div, length={len(result)}")
                return result
            else:
                log("[TextInfo] WYSIWYG textContent was empty")
        except Exception as e:
            log(f"[TextInfo] ERROR reading WYSIWYG: {e}")
    else:
        log("[TextInfo] No WYSIWYG div found")
    log("[TextInfo] Trying broader fallback search in frame...")
    detail_candidates = fr.locator(
        "xpath=("
        "//textarea[contains(@id,'-Text') and (@readonly or @disabled)] | "
        "//*[@role='textbox' and (not(@contenteditable) or @contenteditable='false')] | "
        "//div[contains(@class,'th-wysi') or contains(@class,'th-txt')][not(@contenteditable) or @contenteditable='false'] | "
        "//div[contains(@class,'text-value') or contains(@class,'TextValue')]"
        ")"
    )
    if detail_candidates.count():
        log(f"[TextInfo] Found {detail_candidates.count()} fallback candidates")
        try:
            raw = detail_candidates.first.inner_text()
        except Exception:
            raw = detail_candidates.first.evaluate("n => n.textContent || ''")
        if raw:
            log(f"[TextInfo] ✓ SUCCESS: Got text from fallback search, length={len(raw)}")
            return _normalize_text_preserve(raw) if preserve_format else _normalize_text(raw)
        else:
            log("[TextInfo] Fallback candidates returned empty text")
    else:
        log("[TextInfo] No fallback candidates found")
    log("[TextInfo] Trying last resort: looking for 'Text' label...")
    lab = fr.locator("xpath=//*[normalize-space(.)='Text' or contains(normalize-space(.),'Text')]/following::*[1]").first
    if lab.count():
        log("[TextInfo] Found element following 'Text' label")
        try:
            raw = lab.inner_text()
        except Exception:
            raw = lab.evaluate("n => n.textContent || ''")
        if raw:
            log(f"[TextInfo] ✓ SUCCESS: Got text from 'Text' label follower, length={len(raw)}")
            return _normalize_text_preserve(raw) if preserve_format else _normalize_text(raw)
        else:
            log("[TextInfo] 'Text' label follower returned empty")
    else:
        log("[TextInfo] No 'Text' label found")
    log("[TextInfo] ERROR: All extraction methods failed, returning None")
    return None
def read_text_by_labels_raw(page, wanted_labels):
    fr, tbl = _find_latest_analysis_table_nearby(page)
    if not (fr and tbl and tbl.count()):
        return None
    row = None
    all_rows = tbl.locator(
        "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-TextType')]]"
    )
    row_count = all_rows.count()
    for i in range(row_count):
        candidate_row = all_rows.nth(i)
        type_td = candidate_row.locator(
            "xpath=.//td[starts-with(@id,'GUIDE-TextInfoTable-') and contains(@id,'-TextType')]"
        ).first
        if not type_td.count():
            continue
        type_value = ""
        select = type_td.locator("xpath=.//select").first
        if select.count():
            try:
                selected = select.locator("xpath=.//option[@selected]").first
                if selected.count():
                    type_value = clean(selected.inner_text())
                else:
                    type_value = clean(select.evaluate("""
                        el => {
                            const idx = el.selectedIndex;
                            return idx >= 0 && el.options[idx] ? el.options[idx].text : '';
                        }
                    """))
            except Exception:
                pass
        if not type_value:
            type_value = clean(type_td.inner_text())
        if not type_value:
            continue
        type_lower = type_value.strip().lower()
        for label in wanted_labels:
            label_lower = label.strip().lower()
            if type_lower == label_lower or label_lower in type_lower:
                row = candidate_row
                break
        if row:
            break
    if not row:
        for t in wanted_labels:
            cand = tbl.locator(
                "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') "
                "and contains(@id,'-TextType') and normalize-space(.)=" + xpath_literal(t) + "]]"
            ).first
            if cand.count():
                row = cand
                break
    if not row:
        for t in wanted_labels:
            low = t.lower()
            cand = tbl.locator(
                "xpath=.//tr[td[starts-with(@id,'GUIDE-TextInfoTable-') "
                "and contains(@id,'-TextType') and "
                f"contains(translate(normalize-space(.),'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'), '{low}')]]"
            ).first
            if cand.count():
                row = cand
                break
    if not row:
        return None
    td = row.locator(
        "xpath=.//td[starts-with(@id,'GUIDE-TextInfoTable-') "
        "and contains(@id,'-Text') and not(contains(@id,'-TextType'))]"
    ).first
    if not td.count():
        return None
    a = td.locator("xpath=.//a[contains(@id,'text_table') and contains(@id,'lines')]").first
    if a.count():
        full = (a.get_attribute('title') or a.get_attribute('aria-label') or '').strip()
        if full:
            return full
    try:
        raw = td.evaluate("n => n.textContent || ''") or ""
        if raw.strip():
            return raw.strip()
    except Exception:
        pass
    input_el = td.locator("xpath=.//textarea | .//input").first
    if input_el.count():
        try:
            val = input_el.input_value()
            if val:
                return val.strip()
        except Exception:
            pass
    return None
def read_analysis_raw_text_for_current_pli(page):
    labels = [
        "Analysis Summary",
        "Product Analysis Summary",
        "Analysis/Investigation Summary",
        "Analysis/Investigation conclusion",
        "Analysis/Investigation",
    ]
    raw = (read_text_by_labels_raw(page, labels) or "").strip()
    if raw:
        return raw
    robust = (read_text_by_labels(page, labels, preserve_format=True) or "").strip()
    return robust
def read_analysis_summary_for_current_pli(page):
    labels = [
        "Analysis Summary",
        "Product Analysis Summary",
        "Analysis/Investigation Summary",
        "Analysis/Investigation conclusion",
        "Analysis/Investigation",
    ]
    return (read_text_by_labels(page, labels, preserve_format=True) or "").strip()
def wait_for_search_with_retries(page, s, *, max_attempts=8, probe_period_ms=2000,
                                 reload_between_attempts=True, total_timeout_ms=240000):
    import re, time
    start = time.time()
    aad_host_rx = re.compile(r'(?:^|\.)login\.microsoftonline\.com$', re.I)
    def _try_once():
        try:
            loc, ctx, used = wait_find_in_any_frame(page,
                                                    [s.get('selector')] + (s.get('fallback_selectors', []) or []),
                                                    timeout_ms=2500, poll_ms=150)
            return (loc, ctx, used)
        except Exception:
            return (None, None, None)
    attempt = 0
    while attempt < max_attempts and (time.time() - start) * 1000 < total_timeout_ms:
        attempt += 1
        loc, ctx, used = _try_once()
        if loc:
            return (loc, ctx, used)
        page.wait_for_timeout(probe_period_ms)
        host = ""
        try:
            from urllib.parse import urlparse
            host = urlparse(page.url).hostname or ""
        except Exception:
            pass
        if reload_between_attempts and host and not aad_host_rx.search(host):
            try:
                page.reload(wait_until="load")
            except Exception:
                pass
    raise PWTimeout("[SSO] Search not available after retries")
def get_pa_code_to_id(page):
    nav_fr = _find_leftnav_frame(page)
    if not nav_fr:
        any_found = _scan_pa_anywhere(page, "Product Analysis")
        mapping = {}
        for fr, a, txt, did in any_found:
            code = extract_product_code(txt).upper()
            if code and did:
                mapping[code] = did
        log(f"[PA](fallback-section) code→id mapped: {len(mapping)}")
        return mapping
    _ensure_section_expanded(page, "Product Analysis")
    items = _enumerate_pa_items(nav_fr)
    if not items:
        log("[PA] no anchors inside Product Analysis section; scanning section across frames…")
        any_found = _scan_pa_anywhere(page, "Product Analysis")
        mapping = {}
        for fr, a, txt, did in any_found:
            code = extract_product_code(txt).upper()
            if code and did:
                mapping[code] = did
        log(f"[PA](fallback-section) code→id mapped: {len(mapping)}")
        return mapping
    log(f"[PA] detected {len(items)} anchors in Product Analysis section")
    return {(it["code"] or "").upper(): (it["data_id"] or "") for it in items if it["code"] and it["data_id"]}
def click_associated_transactions_tab(page, root_frame):
    return (
        click_tab_by_text(page, root_frame, "Associated Transactions")
        or click_tab_by_text(page, root_frame, "_ovviewset.do_0012")
    )
def _find_assoc_tx_frame(page):
    for fr in page.frames:
        try:
            if fr.locator("xpath=//div[contains(@id,'_Table_bottom') or contains(@id,'_table_bottom')]").first.count():
                return fr
        except Exception:
            pass
    for fr in page.frames:
        try:
            if fr.get_by_role("button", name=re.compile(r"^\s*(Analysis|Investigation)\s*$", re.I)).first.count():
                return fr
            if fr.locator("xpath=//span[contains(@class,'th-bt-span')][.//b[normalize-space(.)='Analysis' or normalize-space(.)='Investigation']]").first.count():
                return fr
        except Exception:
            pass
    return None
def _row_guess_id_type_status(row):
    txid = _row_txid(row)
    status = _row_status_text(row)
    return txid, status
def _collect_unfiltered(fr, header_tbl, body_tbl, scroll_div):
    if not body_tbl or not body_tbl.count():
        return {"product_analysis": [], "investigation": []}
    _scroll_to_load_all_in_div(fr, body_tbl, scroll_div)
    rows = body_tbl.locator("xpath=.//tr[td]")
    pa, inv = [], []
    for i in range(rows.count()):
        row = rows.nth(i)
        txid, gtype, status = _row_guess_id_type_status(row)
        if not txid or not gtype:
            continue
        is_complete = False
        if status:
            s = status.lower()
            is_complete = ("complete" in s) or (s.strip() == "completed") or ("closed - complete" in s)
        else:
            is_complete = False
        if not is_complete:
            continue
        if gtype == "analysis":
            pa.append(txid)
        elif gtype == "investigation":
            inv.append(txid)
    pa = list(dict.fromkeys(pa))
    inv = list(dict.fromkeys(inv))
    log(f"[AssocTx] (unfiltered) PA complete={len(pa)} INV complete={len(inv)}")
    return {"product_analysis": pa, "investigation": inv}
def _assoc_click_filter(fr, label: str) -> bool:
    try:
        btn = fr.get_by_role("button", name=re.compile(rf"^\s*{re.escape(label)}\s*$", re.I)).first
        if btn.count():
            log(f"[AssocTx] clicking filter via role: {label}")
            robust_click(btn, fr)
            try: fr.wait_for_timeout(300)
            except Exception: pass
            return True
    except Exception:
        pass
    try:
        btn = fr.locator(
            "xpath=//span[contains(@class,'th-bt-span')][.//b[normalize-space(.)="
            f"'{label}']]"
        ).first
        if btn.count():
            log(f"[AssocTx] clicking filter via span/b: {label}")
            robust_click(btn, fr)
            try: fr.wait_for_timeout(300)
            except Exception: pass
            return True
    except Exception:
        pass
    try:
        b = fr.locator(f"xpath=//*[normalize-space(.)='{label}']").first
        if b.count():
            cand = b.locator("xpath=ancestor-or-self::*[self::button or self::span or self::a][1]").first
            if cand.count():
                log(f"[AssocTx] clicking filter via generic text: {label}")
                robust_click(cand, fr)
                try: fr.wait_for_timeout(300)
                except Exception: pass
                return True
    except Exception:
        pass
    log(f"[AssocTx] filter button NOT found: {label}")
    return False
def _hdr_indices_from_any(header_tbl, body_tbl):
    labels = []
    if header_tbl and header_tbl.count():
        cells = header_tbl.locator("xpath=.//thead//th|.//thead//td|.//tr[1]/*")
        for i in range(cells.count()):
            c = cells.nth(i)
            t = clean(c.inner_text())
            if not t:
                t = (c.get_attribute("aria-label") or c.get_attribute("id") or "").strip()
            labels.append((i, t or ""))
    if not labels:
        first = body_tbl.locator("xpath=.//tr[td]").first
        cells = first.locator("xpath=.//th|.//td")
        for i in range(cells.count()):
            c = cells.nth(i)
            t = (
                c.get_attribute("aria-label")
                or c.get_attribute("headers")
                or c.get_attribute("id")
                or clean(c.inner_text())
                or ""
            ).strip()
            labels.append((i, t))
    def find_idx(patterns):
        for i, t in labels:
            if any(re.search(p, t or "", re.I) for p in patterns):
                return i
        return None
    idx_id = find_idx([
        r'(?:trans|txn|transaction|work\s*item).*?(?:id|no|number)'
    ])
    idx_type = find_idx([
        r'(?:trans|txn|transaction|work\s*item|related).*?(?:type|category)'
    ])
    idx_status = find_idx([
        r'(?:status|state)\b'
    ])
    idx_product = find_idx([
        r'\bproduct\b',
        r'\bprod\b'
    ])
    return idx_id, idx_type, idx_status, idx_product
def _scroll_to_load_all_in_div(fr, body_tbl, scroll_div):
    if not body_tbl or not body_tbl.count():
        return
    rows = body_tbl.locator("xpath=.//tr[td]")
    target = scroll_div if (scroll_div and scroll_div.count()) else body_tbl
    try:
        box = body_tbl.bounding_box()
        if box:
            fr.mouse.move(box["x"] + box["width"]/2, box["y"] + min(24, box["height"] - 6))
    except Exception:
        pass
    try:
        target.evaluate("n => { n.tabIndex = 0; n.focus(); }")
    except Exception:
        pass
    prev = -1
    stagnant = 0
    for _ in range(240):
        n = rows.count()
        log(f"[AssocTx] rows visible: {n}")
        if n == prev:
            stagnant += 1
            if stagnant >= 6:
                break
        else:
            stagnant = 0
        prev = n
        try:
            if n > 0:
                rows.nth(n - 1).scroll_into_view_if_needed(timeout=500)
        except Exception:
            pass
        try:
            target.press("PageDown")
        except Exception:
            try: fr.keyboard.press("PageDown")
            except Exception: pass
        try:
            target.press("End")
        except Exception:
            try: fr.keyboard.press("End")
            except Exception: pass
        try:
            fr.mouse.wheel(0, 1800)
        except Exception:
            pass        
        fr.wait_for_timeout(140)
_TRANS_HEADER_RX = re.compile(
    r"(transaction|work\s*item|related|type|category|status|state|number|id|no\b|ref|reference)",
    re.I,
)
_PLI_HEADER_RX = re.compile(
    r"(product|description|lot|sn|serial|qty|quantity|uom|unit)", re.I
)
def _table_header_labels(tbl):
    labels = []
    cells = tbl.locator("xpath=.//thead//th|.//thead//td|.//tr[1]/*")
    for j in range(cells.count()):
        c = cells.nth(j)
        t = clean(c.inner_text()) or (c.get_attribute("aria-label") or c.get_attribute("id") or "")
        t = (t or "").strip()
        if t:
            labels.append(t)
    return labels
def _score_header_labels(labels):
    if not labels:
        return -999
    trans_hits = sum(1 for t in labels if _TRANS_HEADER_RX.search(t))
    pli_hits   = sum(1 for t in labels if _PLI_HEADER_RX.search(t))
    return (trans_hits * 3) - (pli_hits * 4)
def _pick_assoc_grid_table(fr):
    best = (-9999, None, None, None, [])
    bodies = fr.locator(
        "xpath=//div[contains(@id,'_Table_bottom') or contains(@id,'_table_bottom')]"
        "//table[contains(@class,'th-clr-table')]"
    )
    for i in range(bodies.count()):
        body = bodies.nth(i)
        header = body.locator(
            "xpath=ancestor::div[1]/preceding-sibling::div[1]//table[contains(@class,'th-clr-table')]"
        ).first
        labels = _table_header_labels(header if header.count() else body)
        score = _score_header_labels(labels)
        if _is_bad_table(body, labels):
            score -= 5000
        sigs = _table_id_signatures(body)
        if any(_ASSOC_CELL_ID_RX.search(s) for s in sigs):
            score += 8000
        labset = set(labels)
        if "Transaction ID" in labset:
            score += 500
        if "Status" in labset:
            score += 300
        sc = body.locator("xpath=ancestor::div[contains(@style,'overflow')][1]").first
        if sc.count():
            score += 50
        if score > best[0]:
            best = (score, header if header.count() else None, body, sc if sc.count() else None, labels)
    if best[1] is None and best[2] is None:
        any_tbls = fr.locator("xpath=//table[contains(@class,'th-clr-table')]")
        for i in range(any_tbls.count()):
            tb = any_tbls.nth(i)
            labels = _table_header_labels(tb)
            score = _score_header_labels(labels)
            if _is_bad_table(tb, labels): score -= 5000
            sigs = _table_id_signatures(tb)
            if any(_ASSOC_CELL_ID_RX.search(s) for s in sigs): score += 8000
            if score > best[0]:
                best = (score, None, tb, None, labels)
    _, header_tbl, body_tbl, scroll_div, headers = best
    if body_tbl:
        log(f"[AssocTx] chose grid with headers: {headers}")
        if _is_bad_table(body_tbl, headers) and not _is_assoc_tx_table(headers, body_tbl):
            log("[AssocTx] rejecting non-transaction grid (attachments/text/partners)")
            return None, None, None, []
    else:
        log("[AssocTx] no suitable transaction grid found in this frame")
    return header_tbl, body_tbl, scroll_div, headers
def _node_signature(loc):
    try:
        return loc.evaluate("n => (n.innerText || '').slice(0, 1200)") or ""
    except Exception:
        try:
            return clean(loc.inner_text())[:1200]
        except Exception:
            return ""
_ASSOC_CELL_ID_RX = re.compile(r'\bGUIDE-AssociatedTransactionsTable-', re.I)
_BAD_TABLE_ID_RXS = [
    re.compile(r'\bGUIDE-AttachmentsTable-', re.I),
    re.compile(r'\bGUIDE-TextInfoTable-', re.I),
    re.compile(r'\bGUIDE-PartnersTable-', re.I),
    re.compile(r'\bGUIDE-ProductLineItemsTable-', re.I),
]
def _table_id_signatures(tbl):
    sigs = []
    try:
        cells = tbl.locator("xpath=.//th|.//td")
        for j in range(min(cells.count(), 200)):
            sid = (cells.nth(j).get_attribute("id") or "")
            if sid:
                sigs.append(sid)
    except Exception:
        pass
    return sigs
def _is_assoc_tx_table(headers, body_tbl):
    labels = [h.lower() for h in (headers or [])]
    if ("transaction id" in labels and "status" in labels):
        return True
    sigs = _table_id_signatures(body_tbl)
    if any(_ASSOC_CELL_ID_RX.search(s) for s in sigs):
        return True
    return False
def _is_bad_table(body_tbl, headers):
    sigs = _table_id_signatures(body_tbl)
    if any(rx.search(s) for rx in _BAD_TABLE_ID_RXS for s in sigs):
        return True
    h = set((headers or []))
    if {"Name","Document Type","Folder Path"}.issubset(h):
        return True  # attachments
    return False
def read_associated_transactions_complete(page, root_frame):
    click_associated_transactions_tab(page, root_frame)
    fr = None
    for _ in range(20):
        fr = _find_assoc_tx_frame(page)
        if fr:
            break
        page.wait_for_timeout(250)
    if not fr:
        log("[AssocTx] grid frame not found")
        return {"product_analysis": [], "investigation": [], "tx_product_map": {}}
    tx_product_map = {}
    def _collect_for(label):
        header_tbl_0, body_tbl_0, scroll_div_0, headers_0 = _pick_assoc_grid_table(fr)
        sig_before = _node_signature(body_tbl_0) if body_tbl_0 else ""
        if not _assoc_click_filter(fr, label):
            return []
        header_tbl = body_tbl = scroll_div = None
        headers = []
        for _ in range(3):  # ~600ms total
            page.wait_for_timeout(200)
            h1, b1, sc1, hdr1 = _pick_assoc_grid_table(fr)
            if b1 and _is_assoc_tx_table(hdr1, b1):
                sig_after = _node_signature(b1)
                header_tbl, body_tbl, scroll_div, headers = h1, b1, sc1, hdr1
                if (sig_after and sig_after != sig_before) or not sig_before:
                    break
        if not body_tbl:
            log(f"[AssocTx] no AssociatedTransactions grid after clicking {label} — skipping")
            return []
        _scroll_to_load_all_in_div(fr, body_tbl, scroll_div)
        rows = body_tbl.locator("xpath=.//tr[td]")
        idx_id, idx_type, idx_status, idx_product = _hdr_indices_from_any(header_tbl, body_tbl)
        ids_complete, ids_any = [], []
        n = rows.count()
        for i in range(n):
            row = rows.nth(i)
            cells = row.locator("xpath=.//th|.//td")
            txid = None
            if idx_id is not None and cells.count() > idx_id:
                txid_txt = clean(cells.nth(idx_id).inner_text())
                txid = _find_first_match(
                    [r"\b\d{5,}\b", r"\b[A-Z]{2,5}[-_ ]?\d{4,}\b"],
                    txid_txt
                ) or txid_txt
            else:
                txid = _row_txid(row)
            if idx_status is not None and cells.count() > idx_status:
                status = clean(cells.nth(idx_status).inner_text())
            else:
                status = _row_status_text(row)
            product_txt = ""
            if idx_product is not None and cells.count() > idx_product:
                product_txt = clean(cells.nth(idx_product).inner_text())
            if not product_txt:
                prod_cell = row.locator(
                    "xpath=.//td[contains(@id,'GUIDE-AssociatedTransactionsTable') "
                    "and contains(@id,'-Product')]"
                ).first
                if prod_cell.count():
                    a = prod_cell.locator("xpath=.//a").first
                    if a.count():
                        product_txt = clean(
                            a.get_attribute("title")
                            or a.get_attribute("aria-label")
                            or a.inner_text()
                        )
                    else:
                        product_txt = clean(prod_cell.inner_text())
            product_code = extract_product_code(product_txt).upper() if product_txt else ""
            dbg = (
                f"id={txid or '-'} "
                f"status={(status or '').strip()!r} "
                f"product={product_txt!r} code={product_code!r}"
            )
            log(f"[AssocTx:{label}] row {i+1}/{n}: {dbg}")
            if txid and product_code:
                tx_product_map.setdefault(txid, product_code)
            if not txid:
                continue
            ids_any.append(txid)
            if _row_is_complete(row) or re.search(r'\bcomplete(d)?\b', (status or ''), re.I):
                ids_complete.append(txid)
        ids_complete = list(dict.fromkeys(ids_complete))
        ids_any = list(dict.fromkeys(ids_any))
        if ids_complete:
            return ids_complete
        if ids_any:
            log(f"[AssocTx:{label}] no explicit 'Complete' statuses found; returning all ({len(ids_any)})")
            return ids_any
        return []
    has_btns = (
        fr.get_by_role("button", name=re.compile(r"^\s*(Analysis|Investigation)\s*$", re.I)).first.count()
        or fr.locator(
            "xpath=//span[contains(@class,'th-bt-span')][.//b[normalize-space(.)='Analysis' or normalize-space(.)='Investigation']]"
        ).first.count()
    )
    if not has_btns:
        log("[AssocTx] filter buttons not present in detected frame; falling back to unfiltered parse")
        header_tbl, body_tbl, scroll_div, _ = _pick_assoc_grid_table(fr)
        if not body_tbl:
            return {"product_analysis": [], "investigation": [], "tx_product_map": {}}
        base = _collect_unfiltered(fr, header_tbl, body_tbl, scroll_div)
        base["tx_product_map"] = tx_product_map
        return base
    pa = _collect_for("Analysis")
    inv = _collect_for("Investigation")
    log(f"[AssocTx] Product Analysis (Complete or fallback): {pa}")
    log(f"[AssocTx] Investigations (Complete or fallback): {inv}")
    if tx_product_map:
        for txid, code in tx_product_map.items():
            log(f"[AssocTx] TX {txid} → Product {code}")

    return {
        "product_analysis": pa,
        "investigation": inv,
        "tx_product_map": tx_product_map,
    }
IMAGE_ONLY_PA_LEAD_TEMPLATE = (
    "Information provided to Medtronic indicated that the device is not available for evaluation. "
    "However, a picture for the {product_desc} was provided for evaluation. "
    "The review of the provided picture is described below."
)
def summary_has_product_id(text: str, product_id: str) -> bool:
    if not text or not product_id:
        return False
    pid = str(product_id).strip()
    return re.search(rf"\b{re.escape(pid)}\b", text, re.IGNORECASE) is not None
def robust_click(el, frame, timeout_ms=8000):
    try:
        el.scroll_into_view_if_needed(timeout=2000)
    except Exception:
        pass
    try:
        el.click()
        return True
    except Exception:
        pass
    try:
        el.evaluate("e => { e.scrollIntoView({block:'center'}); e.click(); }")
        return True
    except Exception:
        pass
    try:
        box = el.bounding_box()
        if box:
            frame.mouse.click(box["x"] + min(5, box["width"]/2), box["y"] + min(5, box["height"]/2))
            return True
    except Exception:
        pass
    return False
def expand_full_text_if_collapsed(frame):
    try:
        link = frame.locator("xpath=//a[contains(@id,'text_table') and contains(@id,'lines')]").first
        if link.count():
            robust_click(link, frame)
            frame.wait_for_timeout(200)
    except Exception:
        pass
def extract_product_code(desc: str) -> str:
    s = (desc or "").upper()
    toks = re.findall(r"[A-Z0-9][A-Z0-9_-]*", s)
    if not toks:
        return ""
    for t in toks:
        if re.search(r"\d", t):
            return t
    return toks[0]
def active_content_frames(page, cfg):
    name_rx = cfg.get('content_frame_name_regex') or cfg.get('frame_name_regex')
    url_rx  = cfg.get('frame_url_regex')
    frames = []
    for fr in page.frames:
        try:
            nm, url = (fr.name or ""), (fr.url or "")
        except Exception:
            nm, url = "", ""
        if (name_rx and re.search(name_rx, nm, re.I)) or (url_rx and re.search(url_rx, url, re.I)):
            frames.append(fr)
    if frames:
        return frames
    for fr in page.frames:
        if (fr.name or "") == "WorkAreaFrame1":
            return [fr]
    return list(page.frames)
def _find_latest_text_table_in(fr):
    scope = fr.locator("xpath=//div[contains(@class,'th-clr-cnt-bottom')]")
    scope = scope if scope.count() else fr
    tds = scope.locator("xpath=.//table[.//td[starts-with(@id,'GUIDE-TextInfoTable-')]]")
    return tds.nth(tds.count()-1) if tds.count() else None
def _suppress_clicks_enable(frame):
    try:
        frame.evaluate("""
            window.__mdt_suppress = e => { e.stopPropagation(); e.preventDefault(); };
            document.addEventListener('mousedown', window.__mdt_suppress, true);
            document.addEventListener('click', window.__mdt_suppress, true);
            document.addEventListener('mouseup', window.__mdt_suppress, true);
        """)
    except Exception:
        pass
def _suppress_clicks_disable(frame):
    try:
        frame.evaluate("""
            if (window.__mdt_suppress) {
                document.removeEventListener('mousedown', window.__mdt_suppress, true);
                document.removeEventListener('click', window.__mdt_suppress, true);
                document.removeEventListener('mouseup', window.__mdt_suppress, true);
                window.__mdt_suppress = null;
            }
        """)
    except Exception:
        pass
def _normalize_text(s: str) -> str:
    s = (s or "").replace("\xa0", " ")
    s = re.sub(r'\r?\n\s*\r?\n+', '\n\n', s)
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r'\s*\n\s*', '\n', s).strip()
    return s
def _pa_try_expand(fr):
    header = fr.locator("xpath=//*[contains(@class,'left-nav')]//*[contains(@class,'ProductAnalysis')]").first
    if not header.count(): 
        return False
    clicker = header.locator("xpath=following-sibling::*[contains(@class,'clicker')][1]").first
    container = header.locator("xpath=following-sibling::*[contains(@class,'data-wrapper')][1]").first
    if container.count():
        try:
            is_hidden = container.evaluate("n => getComputedStyle(n).display === 'none'")
        except Exception:
            is_hidden = False
        if not is_hidden:
            return True
    if clicker.count():
        try:
            robust_click(clicker, fr)
            fr.wait_for_timeout(250)
        except Exception:
            pass
    if container.count():
        try:
            return not container.evaluate("n => getComputedStyle(n).display === 'none'")
        except Exception:
            return True
    return False
def _pa_anchor_by_data_id(fr, data_id: str):
    if not data_id:
        return None
    return fr.locator(
        "xpath=//*[contains(@class,'left-nav')]"
        "//div[contains(@class,'ProductAnalysis')]"
        "/following-sibling::div[contains(@class,'clicker')][1]"
        "/following-sibling::div[contains(@class,'data-wrapper')][1]"
        f"//a[contains(@class,'GUIDE-sideNav') and @data-trans-id='{data_id}']"
    ).first
def robust_click_plus(el, frame):
    try:
        el.scroll_into_view_if_needed(timeout=1500)
    except Exception:
        pass
    try:
        el.click(timeout=3000, force=True)
        return True
    except Exception:
        pass
    try:
        el.evaluate("""
            e => {
                const ev1 = new MouseEvent('mousedown', {bubbles:true, cancelable:true});
                const ev2 = new MouseEvent('mouseup',   {bubbles:true, cancelable:true});
                const ev3 = new MouseEvent('click',     {bubbles:true, cancelable:true});
                e.dispatchEvent(ev1); e.dispatchEvent(ev2); e.dispatchEvent(ev3);
            }
        """)
        return True
    except Exception:
        pass
    try:
        el.evaluate("e => { e.scrollIntoView({block:'center'}); e.click(); }")
        return True
    except Exception:
        pass
    try:
        box = el.bounding_box()
        if box:
            frame.mouse.click(box["x"] + min(5, box["width"]/2), box["y"] + min(5, box["height"]/2))
            return True
    except Exception:
        pass
    return False
def _list_placeholders(xml: str):
    hits = []
    for m in _PLACEHOLDER_FINDER.finditer(xml):
        i = m.start()
        last_lt = xml.rfind('<', 0, i)
        last_gt = xml.rfind('>', 0, i)
        if last_lt > last_gt:  # we are inside a tag
            seg = xml[last_lt:i]
            if '="' in seg or "='" in seg:
                continue
        label = re.sub(r'<[^>]*?>', '', m.group('label')).strip()
        hits.append(label)
    return hits
def _objects_dropdown_button(fr):
    return fr.locator("xpath=//a[contains(@id,'_Objects-btn') and contains(@class,'th-ip-h')]").first
def _objects_dropdown_list(fr):
    return fr.locator("xpath=//ul[contains(@id,'_Objects_items')]").first
def _click_objects_dropdown(fr):
    btn = _objects_dropdown_button(fr)
    if not btn.count():
        return False
    try:
        robust_click(btn, fr)
        fr.wait_for_timeout(150)
        return True
    except Exception:
        return False
def _list_objects_options(fr):
    ul = _objects_dropdown_list(fr)
    if not ul.count():
        return []
    labels = ul.locator(
        "xpath=.//span[contains(@class,'th-hb-value')] | .//*[@role='option'] | .//a | .//li"
    )  
    out = []
    for i in range(min(labels.count(), 200)):
        el = labels.nth(i)
        try:
            txt = (el.inner_text() or "").strip()
        except Exception:
            txt = (el.get_attribute("aria-label") or el.get_attribute("title") or "").strip()
        if txt:
            out.append((el, txt))
    return out
def _pick_dropdown_option(fr, want_texts) -> bool:
    ul = _objects_dropdown_list(fr)
    if not ul.count():
        return False
    options = _list_objects_options(fr)
    log("[scope] Objects options: " + ", ".join([repr(t) for _, t in options[:30]]))
    wants = [w.strip().lower() for w in (want_texts if isinstance(want_texts, (list, tuple)) else [want_texts])]
    def _match(txt):
        low = (txt or "").strip().lower()
        if any(low == w for w in wants):
            return True
        if any(w in low for w in wants):
            return True
        return "activit" in low
    target = None
    for el, txt in options:
        if _match(txt):
            target = el
            break
    if not target:
        return False
    clickable = target.locator("xpath=ancestor-or-self::*[self::a or self::li or self::div][1]").first
    if not clickable.count():
        clickable = target
    try:
        robust_click(clickable, fr)
        fr.wait_for_timeout(200)
        return True
    except Exception:
        return False
def _objects_button_text(fr) -> str:
    btn = _objects_dropdown_button(fr)
    if not btn.count():
        return ""
    try:
        txt = btn.inner_text().strip()
    except Exception:
        txt = (btn.get_attribute("title") or btn.get_attribute("aria-label") or "").strip()
    return txt
def set_search_scope(page, desired_labels=("Activities","Activity")) -> bool:
    fr = _find_scope_frame_for_objects(page)
    if not fr:
        log("[scope] Could not find frame for Objects dropdown")
        return False
    before = _objects_button_text(fr)
    if not _click_objects_dropdown(fr):
        log("[scope] Could not open Objects dropdown")
        return False
    picked = _pick_dropdown_option(fr, list(desired_labels))
    if not picked:
        log("[scope] No matching option found; keeping scope as-is")
        try: _click_objects_dropdown(fr)
        except Exception: pass
        return False
    fr.wait_for_timeout(150)
    after = _objects_button_text(fr)
    ok = any(lbl.lower() in (after or "").lower() for lbl in desired_labels) or ("activit" in (after or "").lower())
    log(f"[scope] Objects now shows: {after!r} ⇒ ok={ok}")
    return ok
def _find_scope_frame_for_objects(page):
    for fr in page.frames:
        try:
            have = (
                fr.locator("xpath=//*[normalize-space(.)='Search For:']").first.count() and
                fr.locator("xpath=//*[normalize-space(.)='with:']").first.count() and
                (fr.locator("text=Go").first.count() or fr.locator("xpath=//input[@value='Go']|//a[normalize-space(.)='Go']").first.count())
            )
            if have:
                return fr
        except Exception:
            continue
    for fr in page.frames:
        try:
            if _objects_dropdown_button(fr).count() or _objects_input(fr).count():
                return fr
        except Exception:
            continue
    return page.main_frame
def _current_scope_text(fr) -> str:
    val = _read_objects_value(fr)
    if val: return val
    btn = _objects_dropdown_button(fr)
    if btn.count():
        try:
            t = (btn.inner_text() or btn.get_attribute("title") or btn.get_attribute("aria-label") or "").strip()
            if t: return t
        except Exception: pass
    try:
        blk = fr.locator("xpath=//*[normalize-space(.)='Search For:']/following::*[1]").first
        if blk.count():
            t = (blk.inner_text() or blk.get_attribute("title") or blk.get_attribute("aria-label") or "").strip()
            if t: return t
    except Exception:
        pass
    return ""
def set_search_scope_to_activities(page) -> bool:
    fr = _find_scope_frame_for_objects(page)
    if not fr:
        log("[scope] Could not find frame for Objects dropdown")
        return False
    if not _click_objects_dropdown(fr):
        log("[scope] Could not open Objects dropdown")
        return False
    picked = _pick_dropdown_option(fr, "Activities")
    log(f"[scope] Set scope to Activities: {picked}")
    return picked
def _find_global_search_input_in_frame(fr):
    sels = [
        "xpath=//input[contains(@id,'SearchValue') and contains(@class,'th-sif')]",
        "xpath=//input[contains(@id,'SearchValue')]",
        "xpath=//input[contains(@tempname,'search')]",
        "css=input.th-sif",
    ]
    for sel in sels:
        try:
            loc = fr.locator(sel).first
            if loc.count():
                try:
                    loc.wait_for(state="visible", timeout=1500)
                except Exception:
                    pass
                return loc
        except Exception:
            pass
    return None
def _objects_input(fr):
    return fr.locator("xpath=//input[contains(@id,'_Objects') and contains(@class,'th-if') and @role='combobox']").first
def _open_objects_popup(fr):
    inp = _objects_input(fr)
    if inp.count():
        try:
            inp.scroll_into_view_if_needed()
            inp.click()
            fr.wait_for_timeout(80)
            inp.press("ArrowDown")   # opens the list in most builds
            fr.wait_for_timeout(120)
            return True
        except Exception:
            pass
    return _click_objects_dropdown(fr)
def _popup_options_any(fr):
    candidates = []
    ul = fr.locator("xpath=//ul[contains(@id,'_Objects_items') and (not(@style) or not(contains(@style,'display: none')))]").first
    if ul.count():
        items = ul.locator("xpath=.//li | .//a | .//span")
        for i in range(min(items.count(), 200)):
            el = items.nth(i)
            try:
                txt = (el.inner_text() or "").strip()
            except Exception:
                txt = (el.get_attribute("aria-label") or el.get_attribute("title") or "").strip()
            if txt:
                clicky = el.locator("xpath=ancestor-or-self::*[self::li or self::a][1]").first
                candidates.append((clicky if clicky.count() else el, txt))
    listbox = fr.locator("[role='listbox']").first
    if listbox.count():
        opts = listbox.locator("[role='option']")
        for i in range(min(opts.count(), 200)):
            el = opts.nth(i)
            txt = (el.inner_text() or el.get_attribute("aria-label") or "").strip()
            if txt:
                candidates.append((el, txt))
    thpop = fr.locator("xpath=//*[contains(@class,'th-ddlb') or contains(@class,'th-popup')][not(contains(@style,'display: none'))]").first
    if thpop.count():
        opts = thpop.locator("xpath=.//li | .//a | .//span")
        for i in range(min(opts.count(), 200)):
            el = opts.nth(i)
            txt = (el.inner_text() or el.get_attribute("aria-label") or el.get_attribute("title") or "").strip()
            if txt:
                clicky = el.locator("xpath=ancestor-or-self::*[self::li or self::a][1]").first
                candidates.append((clicky if clicky.count() else el, txt))
    return candidates
def _pick_popup_option_by_text(fr, *want_texts):
    wants = [w.strip().lower() for w in want_texts if w]
    if not wants:
        wants = ["activities", "activity"]
    _open_objects_popup(fr)
    fr.wait_for_timeout(120)
    options = _popup_options_any(fr)
    if options:
        log("[scope] Popup options: " + ", ".join(repr(t) for _, t in options[:30]))
    def _match(txt):
        low = (txt or "").lower().strip()
        if any(low == w for w in wants): return True
        if any(w in low for w in wants): return True
        return "activit" in low
    for el, txt in options:
        if _match(txt):
            try:
                robust_click(el, fr)
                fr.wait_for_timeout(120)
                return True
            except Exception:
                pass
    return False
def _read_objects_value(fr) -> str:
    inp = _objects_input(fr)
    if not inp.count(): return ""
    try:
        return (inp.input_value() or "").strip()
    except Exception:
        try:
            return (inp.get_attribute("value") or "").strip()
        except Exception:
            return ""
def _direct_set_objects_value(fr, value: str) -> bool:
    inp = _objects_input(fr)
    if not inp.count():
        return False
    try:
        fr.evaluate(
            """
            (id, val) => {
              const el = document.getElementById(id);
              if (!el) return false;
              const wasRO = el.hasAttribute('readonly');
              if (wasRO) el.removeAttribute('readonly');
              const old = el.value;
              el.value = val;
              el.dispatchEvent(new Event('input', { bubbles: true }));
              el.dispatchEvent(new Event('change', { bubbles: true }));
              try { if (window.thtmlbAutoSave) thtmlbAutoSave(el); } catch (e) {}
              try { if (window.th_ddlb_onchange) th_ddlb_onchange(el); } catch (e) {}
              try { el.blur(); el.focus(); } catch (e) {}
              if (wasRO) el.setAttribute('readonly','readonly');
              return true;
            }
            """,
            inp.get_attribute("id"),
            value,
        )
        fr.wait_for_timeout(120)
        return True
    except Exception:
        return False
def force_scope_to_activities(page) -> bool:
    fr = _find_scope_frame_for_objects(page)
    if not fr:
        log("[scope] header frame not found")
        return False
    cur = (_current_scope_text(fr) or "").lower()
    if "activit" in cur:
        log(f"[scope] already Activities: {cur!r}")
        return True
    if _pick_popup_option_by_text(fr, "Activities", "Activity"):
        fr.wait_for_timeout(160)
        cur = (_current_scope_text(fr) or "").lower()
        if "activit" in cur:
            log(f"[scope] selected via popup: {cur!r}")
            return True
    inp = _objects_input(fr)
    if inp.count():
        try:
            inp.click()
            fr.wait_for_timeout(60)
            try: inp.press("Alt+ArrowDown")
            except Exception: pass
            inp.press("Control+A"); fr.wait_for_timeout(40)
            inp.type("Activities", delay=15); fr.wait_for_timeout(120)
            inp.press("Enter"); fr.wait_for_timeout(180)
            cur = (_current_scope_text(fr) or "").lower()
            if "activit" in cur:
                log("[scope] selected via keyboard type+Enter")
                return True
        except Exception:
            pass
    if _direct_set_objects_value(fr, "Activities"):
        fr.wait_for_timeout(160)
        cur = (_current_scope_text(fr) or "").lower()
        log(f"[scope] after direct-set, now={cur!r}")
        if "activit" in cur:
            return True
    try:
        fr.evaluate("""
            (() => {
              const el = document.querySelector("input[id*='_Objects'][role='combobox']") || document.querySelector("input[id*='_Objects']");
              if (!el) return false;
              const ev = new Event('change', {bubbles:true});
              el.dispatchEvent(ev);
              try { if (window.th_ddlb_onchange) th_ddlb_onchange(el); } catch(e){}
              try { if (window.thtmlbAutoSave) thtmlbAutoSave(el); } catch(e){}
              return true;
            })();
        """)
        fr.wait_for_timeout(160)
        cur = (_current_scope_text(fr) or "").lower()
        if "activit" in cur:
            log("[scope] selected via change hooks")
            return True
    except Exception:
        pass
    log("[scope] FAILED to set Activities (still showing %r)" % _current_scope_text(fr))
    return False
def search_activities_for_id(page, txid: str) -> bool:
    fr = _find_scope_frame_for_objects(page)
    if not fr:
        log("[search] header frame not found")
        return False
    if not force_scope_to_activities(page):
        log("[search] scope not Activities; aborting to avoid SR/PE results")
        return False
    input_el = _find_global_search_input_in_frame(fr)
    if not input_el:
        log("[search] header search input not found")
        return False
    try:
        input_el.click()
        try: input_el.fill("")
        except Exception: pass
        input_el.type(str(txid), delay=18)
        try:
            input_el.evaluate("el => el.dispatchEvent(new Event('input',{bubbles:true}))")
        except Exception: pass
        if not soft_click_go(fr):
            input_el.press("Enter")
        try: fr.wait_for_load_state("networkidle")
        except Exception: pass
        fr.wait_for_timeout(800)
        if "activit" not in (_current_scope_text(fr) or "").lower():
            log("[search] scope reverted after submit; likely clicked content-frame Go. Aborting.")
            return False
        return True
    except Exception as e:
        log(f"[search] error: {e}")
        return False
def _row_has_productid_placeholder(row) -> bool:
    t = _row_text(row)
    return bool(re.search(r'\bproduct\s*id\b', t))
def _table_is_first_products_table(tbl: Table) -> bool:
    for r in tbl.rows:
        if _looks_like_products_header_row(r):
            return True
    return False
def _renumber_xml_placeholders(xml: str, src_idx: int, dst_idx: int) -> str:
    gap = r'(?:\s|<[^>]*?>)*?'  # tolerate tags/runs
    def w(s):  # tolerant words
        return _split_tolerant(s)
    bases = [
        "analysis", "investigation",
        "product id", "product_id",
        "product desc", "product description", "product_desc",
        "lot/serial number", "lot_serial_number",
        "serial no/lot no", "serial_no_lot_no",
        "serial or lot", "serial_or_lot",
    ]
    for base in bases:
        pat = re.compile(
            r'(?P<open>\{\{|\[\[)'+gap+
            w(base)+gap+
            r'(?:[_/]|'+gap+r')?'+gap+str(src_idx)+gap+
            r'(?P<close>\}\}|\]\])',
            re.I|re.S
        )
        xml = pat.sub(lambda m: f"{m.group('open')}{base} {dst_idx}{m.group('close')}", xml)
    return xml
def _find_second_table_and_blocks(doc: Document):
    first_products_tbl_seen = False
    for tbl in doc.tables:
        if _table_is_first_products_table(tbl) and not first_products_tbl_seen:
            first_products_tbl_seen = True
            continue
        starts = [i for i, r in enumerate(tbl.rows) if _row_has_productid_placeholder(r)]
        if not starts:
            continue
        whole_text = " ".join(_row_text(r) for r in tbl.rows)
        maybe_analysisy = ("analysis" in whole_text.lower()) or ("investigation" in whole_text.lower())
        if not maybe_analysisy:
            continue
        blocks = []
        for j, s_idx in enumerate(starts):
            e_idx = (starts[j+1] - 1) if (j + 1 < len(starts)) else (len(tbl.rows) - 1)
            blocks.append((s_idx, e_idx))
        return tbl, blocks
    return None, []
def _ensure_second_table_product_blocks(doc: Document, product_count: int):
    tbl, blocks = _find_second_table_and_blocks(doc)
    if not tbl or not blocks:
        return
    current = len(blocks)
    if product_count == current:
        return
    def _delete_block(tbl: Table, start_idx: int, end_idx: int):
        for i in range(end_idx, start_idx-1, -1):
            _delete_row(tbl, i)
    if product_count < current:
        for b in range(current-1, product_count-1, -1):
            s, e = blocks[b]
            _delete_block(tbl, s, e)
        return
    last_s, last_e = blocks[-1]
    try:
        last_text = _row_text(tbl.rows[last_s])
        m = re.search(r'\bproduct\s*id\s*([0-9]+)\b', last_text)
        src_idx = int(m.group(1)) if m else current
    except Exception:
        src_idx = current
    for new_i in range(current+1, product_count+1):
        for r_idx in range(last_s, last_e+1):
            tr = tbl.rows[r_idx]._tr
            xml = tr.xml
            xml2 = _renumber_xml_placeholders(xml, src_idx, new_i)
            tbl._tbl.append(parse_xml(xml2))
def read_analysis_summary_for_txid(page, txid: str) -> str:
    ok = search_activities_for_id(page, txid)
    if not ok:
        return ""
    click_tab_by_text(page, page.main_frame, "Text Info") or click_tab_by_text(page, page.main_frame, "_ovviewset.do_0006")
    txt = read_analysis_summary_for_current_pli(page)
    if not txt:
        page.wait_for_timeout(500)
        txt = read_analysis_summary_for_current_pli(page)
    return (txt or "").strip()
def read_investigation_summary_for_txid(page, txid: str) -> str:
    ok = search_activities_for_id(page, txid)
    if not ok:
        return ""
    click_tab_by_text(page, page.main_frame, "Text Info") or \
    click_tab_by_text(page, page.main_frame, "_ovviewset.do_0006")
    txt = read_investigation_summary_for_current_pli(page)
    if not txt:
        page.wait_for_timeout(500)
        txt = read_investigation_summary_for_current_pli(page)
    return (txt or "").strip()
def read_investigation_summary_for_current_pli(page):
    labels = [
        "Summary of Investigation",
    ]
    return (read_text_by_labels(page, labels, preserve_format=True) or "").strip()
def _apply_plural_s(xml: str, plural: bool) -> str:
    rx = re.compile(r'(\{\{|\[\[)\s*s\s*(\}\}|\]\])', re.I)
    return rx.sub('s' if plural else '', xml)
_BOILERPLATE_SENTENCES = [
    re.compile(r'This\s+report\s+is\s+based\s+on\s+information\s+provided\s+by\s+[^.!?]*[.!?]?\s*', re.I),
    re.compile(r'Returned\s+Product\s+Analysis\s*\(RPA\)\s*Lab[^.!?]*[.!?]?\s*', re.I),
    re.compile(r'(?:The\s+)?RPA\s+Lab\s+received\s+one[^.!?]*[.!?]?\s*', re.I),
    re.compile(r'This\s+complaint\s+will\s+be\s+used\s+for\s+tracking\s+and\s+trending\s+purposes\.?\s*', re.I),
    re.compile(r'The\s+most\s+likely\s+cause\s+was\s+a\s+component\s+failure\.?\s*', re.I),
    re.compile(r'The\s+failure\s+has\s+been\s+escalated\s+to\s+engineering\s+for\s+further\s+analysis\.?\s*', re.I),
]
def _strip_boilerplate_sentences(s: str) -> str:
    if not s:
        return s
    for rx in _BOILERPLATE_SENTENCES:
        s = rx.sub(' ', s)
    s = re.sub(r'[ \t]{2,}', ' ', s)
    s = re.sub(r'\n{3,}', '\n\n', s)
    return s.strip()
def _strip_analysis_phrases(text: str) -> str:
    if not text:
        return ""
    def _remove_sentence_for_phrase(t: str, phrase: str) -> str:
        pattern = re.compile(
            rf'(^|\s*[\.\?!]\s+|\n+){re.escape(phrase)}[^\.?!]*[\.?!]\s*',
            flags=re.IGNORECASE
        )
        def repl(m: re.Match) -> str:
            return m.group(1)
        return pattern.sub(repl, t)
    for phrase in ("Based on the evidence available", "The root cause of this"):
        text = _remove_sentence_for_phrase(text, phrase)
    text = re.sub(r'[ \t]{2,}', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()
def _is_image_pa(raw_text: str, summary_text: str = "") -> bool:
    text = "\n".join([raw_text or "", summary_text or ""])
    if not text.strip():
        return False
    t = re.sub(r"\s+", " ", text).lower()
    has_media = bool(re.search(
        r"\b(picture|photo|image|video|photo\(s\)|picture/video)\b", t, re.I
    ))
    has_provided = bool(re.search(
        r"\b(provided|submitted|reviewed|returned)\b", t, re.I
    ))
    has_no_sample = bool(re.search(
        r"\b(product|device|sample)\s+(was\s+)?not\s+(returned|available)\b", t, re.I
    ))
    has_visual_photo_phrase = bool(re.search(
        r"\bvisual inspection\b.*\b(photo|picture|image)\b|\breturned photo\(s\)\b", t, re.I
    ))
    return (has_media and (has_provided or has_visual_photo_phrase) and has_no_sample) or \
           (has_visual_photo_phrase and has_media)
def _clean_image_pa_text(text: str) -> str:
    if not text:
        return text
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(
        r'(?im)\bthe\s+product\s+sample\s+was\s+not\s+returned\s+to\s+the\s+product\s+analysis\s+laboratory\s*;?\s*',
        '',
        text
    )
    text = re.sub(
        r'^\s*The\s+product\s+sample\s+was\s+not\s+returned\s+to\s+the\s*',
        '',
        text,
        flags=re.IGNORECASE
    ).lstrip()
    text = re.sub(
        r'(?ims)^\s*Evaluation\s*:\s*\n'
        r'(?:\s*(?:[-\u2022\u2013*]|\d+[.)])\s*.*(?:\n|$))*',
        '',
        text
    )
    text = re.sub(
        r'(?im)^\s*Visual\s+inspection\s*:',
        'Picture Evaluation:',
        text
    )
    text = re.sub(
        r'(?ims)\n*\s*A\s+visual\s+inspection\s+of\s+the\s+returned\s+photo\(s\)\s+noted\s*:\s*\n?',
        '\n',
        text
    )
    text = re.sub(
        r'(?im)\bProduct\s+Analysis\s+laboratory\s*;\s*however,\s*'
        r'a\s+picture\s*/\s*video\s+was\s+provided\s+by\s+the\s+customer\s+for\s+analysis\.?\s*',
        '',
        text
    )
    text = re.sub(
        r'(?im)\bProduct\s+Analysis\s+laboratory\s*;\s*however,\s*'
        r'a\s+picture\s+was\s+provided\s+by\s+the\s+customer\s+for\s+analysis\.?\s*',
        '',
        text
    )
    text = re.sub(r'(?m)^\s*-\s+', '\u2022 ', text)
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    text = _strip_media_provided_sentence(text)
    return text
def get_partners_for_ui(frame):
    tbl = _partners_table(frame)
    if not tbl:
        print("[Partners] No partners table found for UI.")
        log("[Partners] No partners table found for UI.")
        return []
    rows = tbl.locator(
        "xpath=.//tr[td[starts-with(@id,'GUIDE-PartnersTable-') and (contains(@id,'-PartnerFunction') or contains(@id,'-Function'))]]"
    )
    if rows.count() == 0:
        log("[Partners] No rows found with PartnerFunction ID, trying all rows with GUIDE-PartnersTable cells")
        rows = tbl.locator(
            "xpath=.//tr[td[starts-with(@id,'GUIDE-PartnersTable-')]]"
        )
    partners = []
    n = rows.count()
    print(f"[Partners] Partner table rows detected: {n}")
    log(f"[Partners] Partner table rows detected: {n}")
    for i in range(n):
        row = rows.nth(i)
        if i == 0:
            cells = row.locator("xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-')]")
            cell_ids = []
            for j in range(min(cells.count(), 5)):
                cell_id = cells.nth(j).get_attribute("id") or ""
                cell_ids.append(cell_id)
            log(f"[Partners] Sample cell IDs from first row: {cell_ids}")
        pf_text = _cell_text_in_same_row(row, "PartnerFunction")
        name = _cell_text_in_same_row(row, "Name")
        addr = _cell_text_in_same_row(row, "Address")
        if not addr:
            addr = _cell_text_in_same_row(row, "address_short")
        if not pf_text:
            pf_cell = row.locator(
                "xpath=.//td[starts-with(@id,'GUIDE-PartnersTable-')]" +
                "[.//select or contains(@id,'Function') or contains(@id,'-Fct')]"
            ).first
            if pf_cell.count():
                select = pf_cell.locator("xpath=.//select").first
                if select.count():
                    try:
                        selected_option = select.locator("xpath=.//option[@selected]").first
                        if selected_option.count():
                            pf_text = clean(selected_option.inner_text())
                        else:
                            try:
                                pf_text = clean(select.evaluate("""
                                    el => {
                                        const idx = el.selectedIndex;
                                        return idx >= 0 && el.options[idx] ? el.options[idx].text : '';
                                    }
                                """))
                            except Exception:
                                pass
                    except Exception as e:
                        log(f"[Partners] Error reading select dropdown: {e}")
                if not pf_text:
                    pf_text = clean(pf_cell.inner_text())
        if addr:
            addr = addr.replace(" / ", "\n")
        log(f"[Partners] Row {i+1}: pf={pf_text!r}, name={name!r}, addr={addr!r}")
        block = "\n".join(x for x in [name, addr] if x).strip()
        partners.append({
            "partner_function": pf_text or "(unknown)",
            "name": name,
            "address": addr,
            "display": (
                f"{pf_text}: {block}" if block and pf_text else
                block or name or pf_text or "(no data)"
            ),
        })
    print(f"[Partners] partners_for_ui count: {len(partners)}")
    log(f"[Partners] partners_for_ui count: {len(partners)}")
    for p in partners:
        print("   -", p["display"].replace("\n", " / "))
    return partners
def build_recipient_options(values: dict):
    partners = values.get("_external_contacts") or []
    name_opts = []
    seen = set()
    for p in partners:
        nm = (p.get("name") or "").strip()
        if not nm:
            continue
        key = nm.lower()
        if key not in seen:
            seen.add(key)
            name_opts.append(nm)
    addr_opts = []
    seen = set()
    for p in partners:
        addr = (p.get("address") or "").strip()
        if not addr:
            continue
        key = addr.lower()
        if key not in seen:
            seen.add(key)
            addr_opts.append(addr)
    default_name = (values.get("ir_name") or "").strip()
    default_addr = ""
    facility_block = (values.get("ir_with_address") or "").strip()
    if facility_block:
        lines = [ln.strip() for ln in facility_block.splitlines() if ln.strip()]
        if len(lines) >= 2:
            default_addr = "\n".join(lines[1:]).strip()
    if default_name and default_name.lower() not in {x.lower() for x in name_opts}:
        name_opts.insert(0, default_name)
    if default_addr and default_addr.lower() not in {x.lower() for x in addr_opts}:
        addr_opts.insert(0, default_addr)
    return {
        "name_options": name_opts,
        "address_options": addr_opts,
        "default_name": default_name,
        "default_address": default_addr,
    }
def _strip_media_provided_sentence(text: str) -> str:
    if not text:
        return text
    text = re.sub(
        r'''(?imx)
        (?:^|[\.\!\?]\s+|\n+)              # sentence boundary
        (?:however,\s*)?                   # optional however
        a\s+
        (?:picture|photo|image)
        (?:\s*/\s*video)?                  # optional "/video"
        \s+was\s+provided\s+by\s+the\s+customer\s+for\s+analysis
        \.?
        (?=\s|$)
        ''',
        ' ',
        text
    )
    text = re.sub(
        r'''(?imx)
        (?:^|[\.\!\?]\s+|\n+)
        (?:however,\s*)?
        (?:pictures|photos|images|videos)\s+were\s+provided\s+by\s+the\s+customer\s+for\s+analysis
        \.?
        (?=\s|$)
        ''',
        ' ',
        text
    )
    text = re.sub(r'[ \t]{2,}', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()
def scrape_complaint(complaint_id: str, cfg_path: str):
    log_filename = setup_logging()
    log(f"Starting scrape for complaint: {complaint_id}")
    log(f"Log file: {log_filename}")
    cfg_path = Path(cfg_path)
    cfg = yaml.safe_load(cfg_path.read_text())
    template_path = Path(cfg['template_path']).expanduser()
    out_dir = Path(cfg.get('output_dir', '.')).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    values = {}
    products = []
    with sync_playwright() as p:
        context, page = launch_gch_context(p, cfg)
        log(f"[browser] Using persistent profile: {resolve_user_data_dir(cfg)}")
        log(f"Navigating to CRM: {cfg['crm_url']}")
        page.goto(cfg['crm_url'], wait_until="load")
        sso_wait = cfg.get('sso_pause_seconds', 0)
        if sso_wait > 0:
            s = cfg.get('search', {})
            try:
                wait_find_in_any_frame(page, [s.get('selector')] + (s.get('fallback_selectors', []) or []),
                                    timeout_ms=2500, poll_ms=150)
                print("[SSO] Search is already available; skipping SSO wait.")
            except Exception:
                print(f"[SSO] Search not ready; running extended SSO retries (up to ~{cfg.get('sso_total_timeout_ms', 240000)//1000}s)…")
                try:
                    wait_for_search_with_retries(
                        page, s,
                        max_attempts=cfg.get('sso_max_attempts', 8),
                        probe_period_ms=cfg.get('sso_probe_period_ms', 2000),
                        reload_between_attempts=cfg.get('sso_reload_between_attempts', True),
                        total_timeout_ms=cfg.get('sso_total_timeout_ms', 240000),
                    )
                except Exception:
                    page.wait_for_timeout(sso_wait * 1000)
        frame = find_app_frame(
            page,
            frame_name_regex=cfg.get('frame_name_regex'),
            url_regex=cfg.get('frame_url_regex')
        )
        if 'search' in cfg:
            s = cfg['search']
            if s.get('selector'):
                try:
                    fallbacks = s.get('fallback_selectors', [])
                    extra_defaults = [
                        "xpath=//input[@id='C12_W37_V38_SearchValue']",
                        "css=#C1_W1_V2_C9_W28_V29_C12_W37_V38_launcher\\.do > span > table > tbody > tr > td > span.th-if-wrapper > input",
                        "xpath=/html/body/form/div[5]/div/table/tbody/tr[1]/td/div/div/div/div/table/tbody/tr/td[1]/div/div/span/table/tbody/tr/td/span[3]/input",
                    ]
                    all_selectors = [s['selector']] + [sel for sel in fallbacks + extra_defaults if sel not in fallbacks]
                    log("[search] polling every 2s for search input…")
                    poll_deadline = time.time() + s.get('pre_wait_timeout', 30000)/1000.0
                    target = target_ctx = used_sel = None
                    while time.time() < poll_deadline:
                        try:
                            target, target_ctx, used_sel = wait_find_in_any_frame(page, all_selectors, timeout_ms=1, poll_ms=1)
                            break
                        except Exception:
                            log("[search] not visible yet; sleeping 2s")
                            time.sleep(2)
                    if not target:
                        target, target_ctx, used_sel = wait_find_in_any_frame(page, all_selectors, timeout_ms=5000)
                    log(f"[search] Found input via selector: {used_sel} in frame url={getattr(target_ctx, 'url', '')} name={getattr(target_ctx, 'name', '')}")
                    if s.get('clear', True):
                        try:
                            target.fill("")
                        except Exception:
                            pass
                    target.click()
                    frame.wait_for_timeout(100)
                    try:
                        log(f"[search] filling complaint id: {complaint_id}")
                        target.fill(complaint_id)
                    except:
                        target.type(complaint_id, delay=30)
                    try:
                        target.evaluate("el => { el.dispatchEvent(new Event('input',{bubbles:true})); el.blur(); el.focus(); }")
                    except: pass
                    submitted = False
                    if s.get('submit_selector'):
                        try:
                            btn = target_ctx.locator(s['submit_selector']).first
                            if btn.count():
                                log("[search] clicking submit")
                                btn.click()
                                submitted = True
                        except Exception:
                            pass
                    if not submitted:
                        key = s.get('press_key', 'Enter')
                        try:
                            log(f"[search] pressing key: {key}")
                            target.press(key)
                        except Exception: pass
                    if s.get('wait_for'):
                        log(f"[wait] waiting for results: {s['wait_for']}")
                        target_ctx.wait_for_selector(s['wait_for'], timeout=s.get('wait_timeout', 60000))
                    else:
                        target_ctx.wait_for_load_state("networkidle")
                        target_ctx.wait_for_timeout(s.get('post_wait_ms', 3000))
                except Exception as e:
                    log(f"[ERROR] Search failed: {e}")
                    print(f"[Search] Failed to drive search: {e}")
                    dump_frames_debug(page, basename='debug')
                    try:
                        page.screenshot(path="debug_search_failure.png", full_page=True)
                        print("Saved debug_search_failure.png")
                        log("Saved debug_search_failure.png")
                    except Exception:
                        pass
        values = {}
        if not values.get("todays_date"):
            values["todays_date"] = datetime.now().strftime("%B %d, %Y").replace(" 0", " ")
        values['complaint_id'] = complaint_id
        for key, conf in cfg.get('field_map', {}).items():
            if isinstance(conf, str):
                values[key] = get_by_label(frame, conf)
            elif isinstance(conf, list):
                values[key] = get_by_label(frame, conf)
            elif isinstance(conf, dict):
                if conf.get('type') == 'label':
                    values[key] = get_by_label(frame, conf.get('labels', []))
                elif conf.get('type') == 'grid':
                    grid_xpath = conf['grid_xpath']
                    row = conf['row']
                    col = conf['col']
                    values[key] = get_grid_row_col(frame, grid_xpath, row, col)
                elif conf.get('type') == 'literal':
                    values[key] = conf.get('value', '')
            else:
                values[key] = ""
        for k, v in cfg.get('defaults', {}).items():
            values.setdefault(k, v)
        try:
            log("[step 1] Partners tab → IR name & facility")
            if click_partners_tab(page, frame):
                pframe = find_partners_frame(page)
                if not pframe:
                    debug_frames_for_partners(page)
                    print("[Partners] Could not locate the partners frame.")
                    log("[Partners] Could not locate the partners frame.")
                else:
                    _debug_list_pf_from_correct_table(pframe)
                    irname = get_initial_reporter_name(pframe)
                    if irname:
                        values['ir_name'] = irname
                    facility_block = get_facility_name_and_address(pframe)
                    if facility_block:
                        values['ir_with_address'] = facility_block
                        lines = [ln.strip() for ln in facility_block.splitlines() if ln.strip()]
                        facility_name = lines[0] if lines else ""
                        facility_address = "\n".join(lines[1:]) if len(lines) > 1 else ""
                        values["facility_name"] = facility_name
                        values["facility_address"] = facility_address
                    partners_for_ui = get_partners_for_ui(pframe)
                    if partners_for_ui:
                        values["_external_contacts"] = partners_for_ui
                        ui_opts = build_recipient_options(values)
                        values["_ui_name_options"] = ui_opts["name_options"]
                        values["_ui_addr_options"] = ui_opts["address_options"]
                        values["_ui_default_name"] = ui_opts["default_name"]
                        values["_ui_default_addr"] = ui_opts["default_address"]
                    print("[Partners] ir_name =", values.get('ir_name', ''))
                    print("[Partners] ir_with_address =", values.get('ir_with_address', ''))
                    log(f"[Partners] ir_name = {values.get('ir_name','')}")
                    log(f"[Partners] ir_with_address = {values.get('ir_with_address','')}")
            else:
                print("[Partners] Could not open Partners tab; leaving ir_* fields from label map/fallbacks.")
                log("[Partners] Could not open Partners tab; leaving ir_* fields from label map/fallbacks.")
        except Exception as e:
            print(f"[Partners] Error scraping Partners tab: {e}")
            log(f"[ERROR] Partners scrape: {e}")
        log("[step 2] Additional External References → rb_reference & report_number")
        ext = read_external_refs(page, frame)
        if ext.get("rb_reference"):
            values["rb_reference"] = ext["rb_reference"]
        if ext.get("report_number"):
            values["report_number"] = ext["report_number"]
        if ext.get("ex_ref"):
            values["ex_ref"] = ext["ex_ref"]
        contacts = ext.get("external_contacts") or []
        values["_aer_external_contacts"] = contacts  
        if contacts and not values.get("external_contact"):
            values["external_contact"] = contacts[0].get("number", "")
        log(
            f"[AER] rb_reference={values.get('rb_reference','')}, "
            f"report_number={values.get('report_number','')}, "
            f"external_contacts={len(contacts)}"
        )
        log("[step 3] Dates tab → event_date")
        event_date_text = get_event_date(page)
        if event_date_text:
            values['event_date'] = event_date_text
        log(f"[Dates] event_date={values.get('event_date','')}")
        log("[step 4] Product Line Items → all rows")
        products = read_all_products(page, frame)
        log(f"[PLI] rows detected: {len(products)}")
        log("[step 5] Text Info → event_description")
        prev_sig = _textinfo_signature(page)
        desc = read_event_description(page, frame)
        if not desc and wait_for_textinfo_change(page, prev_sig, timeout=8000):
            desc = read_event_description(page, frame)
        if desc:
            desc = re.sub(r'^\s*according\s+to\s+the\s+reporter[,:-]?\s*', '', desc, flags=re.I)
            desc = re.sub(r'^\s*it\s+was\s+reported(?:\s+that)?[,:-]?\s*', '', desc, flags=re.I).lstrip()
            desc = re.sub(r'([.!?])\1+', r'\1', desc)
            desc = _remove_product_desc_from_event(desc, products)
            values["event_description"] = desc
        log(f"[Text] description length: {len(values.get('event_description',''))}")
        log("[step 6] Associated Transactions → collect Complete Investigation/Product Analysis IDs")
        assoc = read_associated_transactions_complete(page, frame)
        tx_product_map = assoc.get("tx_product_map", {}) or {}
        code_to_idx = {}
        def _match_summary_to_product_index(
            txid: str,
            summary: str,
            prod_code_from_bc: str,
            products: list,
            code_to_idx: dict,
            tx_product_map: dict,
        ):
            def _fuzzy_lookup(code: str):
                if not code:
                    return None
                c = code.strip().upper()
                if not c:
                    return None
                idx = code_to_idx.get(c)
                if idx:
                    return idx
                candidates = []
                for key, idx in code_to_idx.items():
                    if c in key or key in c:
                        candidates.append((key, idx))
                if len(candidates) == 1:
                    key, idx = candidates[0]
                    log(
                        f"[Match] txid={txid} fuzzy-matched code={c!r} "
                        f"to key={key!r} → product index {idx}"
                    )
                    return idx
                if candidates:
                    log(
                        f"[Match] txid={txid} ambiguous fuzzy match for code={c!r} "
                        f"candidates={[k for k, _ in candidates]}"
                    )
                return None
            assoc_code_raw = tx_product_map.get(txid) or ""
            assoc_idx = _fuzzy_lookup(assoc_code_raw)
            if assoc_idx:
                log(
                    f"[Match] txid={txid} matched via AssocTx grid "
                    f"product={assoc_code_raw.strip().upper()!r} → product index {assoc_idx}"
                )
                return assoc_idx
            bc_code_raw = prod_code_from_bc or ""
            bc_idx = _fuzzy_lookup(bc_code_raw)
            if bc_idx:
                log(
                    f"[Match] txid={txid} matched via bcTitle product="
                    f"{bc_code_raw.strip().upper()!r} → product index {bc_idx}"
                )
                return bc_idx
            for i, p in enumerate(products, start=1):
                pid = (p.get("id") or "").strip()
                if pid and summary_has_product_id(summary, pid):
                    log(
                        f"[Match] txid={txid} matched via summary_has_product_id({pid!r}) "
                        f"→ product index {i}"
                    )
                    return i
            for i, p in enumerate(products, start=1):
                desc = (p.get("desc") or "").strip()
                if not desc:
                    continue
                code_token = extract_product_code(desc)
                if not code_token:
                    continue
                if re.search(rf"\b{re.escape(code_token)}\b", summary, re.I):
                    log(
                        f"[Match] txid={txid} matched via summary product code token "
                        f"{code_token!r} → product index {i}"
                    )
                    return i
            log(f"[Match] txid={txid} could NOT be matched to any product")
            return None
        for idx, p in enumerate(products, start=1):
            pid  = (p.get("id")   or "").strip().upper()
            code = (p.get("code") or extract_product_code(p.get("desc", ""))).upper()
            if pid:
                code_to_idx.setdefault(pid, idx)
            if code:
                code_to_idx.setdefault(code, idx)
        default_pa_text = (
            "Information provided to Medtronic indicated that the complaint device "
            "was not available for evaluation."
        )
        pa_ids_raw = values.get("assoc_tx_product_analysis_ids", "") or ", ".join(assoc.get("product_analysis", []))
        pa_ids = [x.strip() for x in pa_ids_raw.split(",") if x.strip()]
        per_product_pa = {}
        per_product_pa_image = {}
        unmatched_pa = []
        for txid in pa_ids:
            log(f"[PA-SUMMARY] Fetching Analysis Summary for PA ID: {txid}")
            raw_summary, prod_code, raw_txt = read_analysis_summary_and_product_for_txid(page, txid)
            if not (raw_summary or "").strip() and (raw_txt or "").strip():
                log(f"[PA-SUMMARY] txid={txid} summary empty; falling back to raw_txt")
                raw_summary = raw_txt
            is_image = _is_image_pa(raw_txt, raw_summary)
            log(f"[PA-IMG] txid={txid} is_image={is_image} raw_len={len(raw_txt or '')} summary_len={len(raw_summary or '')}")
            if not is_image:
                log(f"[PA-IMG] raw_head={(raw_txt or '')[:200]!r}")
                log(f"[PA-IMG] sum_head={(raw_summary or '')[:200]!r}")
            summary = _normalize_text_preserve(raw_summary)
            summary = _strip_analysis_phrases(summary)
            summary = _strip_media_provided_sentence(summary)
            if not summary:
                summary = "(No Analysis Summary found)"
            idx = _match_summary_to_product_index(
                txid,
                summary,
                prod_code,
                products,
                code_to_idx,
                tx_product_map,
            )
            if idx:
                prod_desc = ""
                if 1 <= idx <= len(products):
                    prod_desc = products[idx - 1].get("desc", "")
                same_product_count = 1
                if 1 <= idx <= len(products):
                    target_id = (products[idx - 1].get("id") or "").strip().upper()
                    target_desc = (products[idx - 1].get("desc") or "").strip().upper()
                    same_product_count = sum(
                        1 for p in products
                        if ((p.get("id") or "").strip().upper() == target_id and target_id)
                        or ((p.get("desc") or "").strip().upper() == target_desc and target_desc)
                    )
                    if same_product_count < 1:
                        same_product_count = 1
                if is_image:
                    cleaned = _clean_image_pa_text(summary)
                    prev_img = per_product_pa_image.get(idx, "")
                    per_product_pa_image[idx] = (
                        prev_img + ("\n\n" if prev_img else "") + cleaned
                    ).strip()
                    log(f"[PA-SUMMARY] txid={txid} stored as IMAGE PA "
                        f"for product index {idx}")
                    log(f"[PA-IMG] appended image block for product idx={idx}; _pa_image_indices will include it")
                else:
                    is_first = idx not in per_product_pa
                    formatted = _format_analysis_block(
                        prod_desc, summary,
                        product_count=same_product_count,
                        include_lead=is_first,
                        has_image=(idx in per_product_pa_image),
                    )
                    prev = per_product_pa.get(idx, "")
                    per_product_pa[idx] = (
                        prev + ("\n\n" if prev else "") + formatted
                    ).strip()
            else:
                unmatched_pa.append(summary)
                log(f"[PA-SUMMARY] txid={txid} could not be matched to any included product; analysis ignored.")
        for idx, img_text in per_product_pa_image.items():
            if idx in per_product_pa:
                existing = per_product_pa[idx]
                existing = re.sub(
                    r'(\bOne\s+.+?)\s+was\s+received\s+for\s+evaluation\.\s*'
                    r'Examination\s+of\s+the\s+sample\s+is\s+described\s+below\.',
                    r'\1 and one picture were received for evaluation. '
                    r'Examination of the sample and picture is provided below.',
                    existing, count=1
                )
                existing = re.sub(
                    r'(\b(?:Two|Three|Four|Five|Six|Seven|Eight|Nine|Ten|\d+)\s+.+?)\s+were\s+received\s+for\s+evaluation\.\s*'
                    r'Examination\s+of\s+the\s+samples?\s+is\s+described\s+below\.',
                    r'\1 and one picture were received for evaluation. '
                    r'Examination of the samples and picture is provided below.',
                    existing, count=1
                )
                insert_done = False
                m = re.search(r'(?im)^\s*Visual inspection\s*:', existing)
                if m:
                    per_product_pa[idx] = (
                        existing[:m.start()].rstrip() + "\n\n" +
                        img_text.strip() + "\n\n" +
                        existing[m.start():].lstrip()
                    )
                    insert_done = True
                if not insert_done:
                    m = re.search(r'(?im)^\s*Evaluation\s*:', existing)
                    if m:
                        per_product_pa[idx] = (
                            existing[:m.start()].rstrip() + "\n\n" +
                            img_text.strip() + "\n\n" +
                            existing[m.start():].lstrip()
                        )
                        insert_done = True
                if not insert_done:
                    per_product_pa[idx] = existing.rstrip() + "\n\n" + img_text.strip()
            else:
                prod_desc = ""
                if 1 <= idx <= len(products):
                    prod_desc = products[idx - 1].get("desc", "")
                desc = (prod_desc or "product").strip()
                lead = IMAGE_ONLY_PA_LEAD_TEMPLATE.format(product_desc=desc)
                per_product_pa[idx] = lead + "\n\n" + img_text.strip()
            log(f"[PA-IMAGE] Product {idx}: inserted image PA text into analysis_{idx}")
        for idx, text in per_product_pa.items():
            values[f"analysis_{idx}"] = text
        values["_pa_image_indices"] = sorted(per_product_pa_image.keys())
        all_pa_blocks = list(per_product_pa.values())
        if all_pa_blocks:
            values["analysis_results"] = "\n\n".join(all_pa_blocks)
        else:
            values["analysis_results"] = default_pa_text
        if len(products) == 1:
            if values.get("analysis_results"):
                values.setdefault("analysis_1", values["analysis_results"])
            else:
                values.setdefault("analysis_1", default_pa_text)
        inv_ids_raw = values.get("assoc_tx_investigation_ids", "") or ", ".join(assoc.get("investigation", []))
        inv_ids = [x.strip() for x in inv_ids_raw.split(",") if x.strip()]
        per_product_inv = {}
        unmatched_inv = []
        for txid in inv_ids:
            log(f"[INV-SUMMARY] Fetching Investigation Summary for INV ID: {txid}")
            raw_summary, prod_code = read_investigation_summary_and_product_for_txid(page, txid)
            summary = _normalize_text_preserve(raw_summary)
            summary = _postprocess_investigation_text(summary)
            if not summary:
                summary = ""
            idx = _match_summary_to_product_index(
                txid,
                summary,
                prod_code,
                products,
                code_to_idx,
                tx_product_map,
            )
            if idx:
                prev = per_product_inv.get(idx, "")
                per_product_inv[idx] = (prev + ("\n\n" if prev else "") + summary).strip()
            else:
                unmatched_inv.append(summary)
                log(f"[INV-SUMMARY] txid={txid} could not be matched to any included product; investigation ignored.")
        for idx, text in per_product_inv.items():
            values[f"investigation_{idx}"] = text
        all_inv_blocks = list(per_product_inv.values())
        if all_inv_blocks:
            values["investigation_summary"] = "\n\n\n".join(all_inv_blocks)
        else:
            values.setdefault("investigation_summary", DEFAULT_INV_TEXT)
        if values.get("analysis_results"):
            values.setdefault("analysis_1", values["analysis_results"])
        else:
            values.setdefault(
                "analysis_1",
                "Information provided to Medtronic indicated that the complaint device "
                "was not available for evaluation."
            )
        if len(products) == 1:
            if values.get("investigation_summary"):
                values.setdefault("investigation_1", values["investigation_summary"])
            else:
                values.setdefault("investigation_1", DEFAULT_INV_TEXT)
        for idx, p in enumerate(products, start=1):
            code = (p.get("code") or extract_product_code(p.get("desc",""))).upper()
            values[f"product_id_{idx}"]   = (p.get("id") or code)
            values[f"product_desc_{idx}"] = p.get("desc", "")
            sn  = _clean_sn(p.get("sn", ""))
            lot = _clean_lot(p.get("lot", ""))
            values[f"product_sn_{idx}"]  = sn
            values[f"product_lot_{idx}"] = lot
            if sn and lot:
                first_table_display = f"SN: {sn} / LN: {lot}"
            elif sn:
                first_table_display = f"SN: {sn}"
            elif lot:
                first_table_display = f"LN: {lot}"
            else:
                first_table_display = ""
            values[f"serial_or_lot_{idx}"] = first_table_display
            if sn and not lot:
                label_2 = "Serial No:"
                value_2 = sn
            elif lot and not sn:
                label_2 = "Lot No:"
                value_2 = lot
            elif sn and lot:
                label_2 = "Serial/Lot No:"
                value_2 = f"{sn} / {lot}"
            else:
                label_2 = ""
                value_2 = ""
            values[f"serial_or_lot_label_{idx}"] = label_2
            values[f"serial_or_lot_value_{idx}"] = value_2
        values["assoc_tx_product_analysis_ids"] = ", ".join(assoc["product_analysis"])
        values["assoc_tx_investigation_ids"]    = ", ".join(assoc["investigation"])
        if any(values.get(f"investigation_{i+1}", "").strip() for i in range(len(products))):
            values["investigation_summary"] = "\n\n\n".join(
                (values.get(f"investigation_{i+1}", "") or "").strip()
                for i in range(len(products))
                if (values.get(f"investigation_{i+1}", "") or "").strip()
            )
        values['_product_count'] = len(products)
        if not (values.get('ir_name') or '').strip():
            values['ir_name'] = 'Customer'
        if len(products) > 3:
            extras = [f"{p['id']} — {p['desc']}" for p in products[3:]]
            values.setdefault("product_extras", "\n".join(extras))
        default_pa_text = (
            "Information provided to Medtronic indicated that the complaint device "
            "was not available for evaluation."
        )
        for idx in range(1, len(products) + 1):
            a_key = f"analysis_{idx}"
            i_key = f"investigation_{idx}"
            analysis_text = (values.get(a_key) or "").strip()
            if not analysis_text or analysis_text == DEFAULT_PA_TEXT:
                values[a_key] = DEFAULT_PA_TEXT
                values[i_key] = DEFAULT_INV_TEXT
            else:
                inv_text = (values.get(i_key) or "").strip()
                if not inv_text:
                    values[i_key] = DEFAULT_INV_TEXT
        for k, v in list(values.items()):
            if not re.match(r"^investigation_\d+$", k):
                continue
            if not (v or "").strip():
                continue
            body = _extract_investigation_body(v)
            values[k] = _strip_leading_based_on_evidence(body)
        log("Collected fields:")
        log(json.dumps(values, indent=2))
        release_input = getattr(context, "release_input", None)
        if callable(release_input):
            release_input()
            log("[browser] Edge window input re-enabled for user editing.")
        log("GCH automation complete; leaving the browser window open for the user.")
    return values, products, cfg, template_path, out_dir
def main():
    if len(sys.argv) < 3:
        print("Usage: python scrape_and_generate.py <complaint_id> <config.yaml>")
        sys.exit(2)
    complaint_id = sys.argv[1]
    cfg_path = sys.argv[2]
    values, products, cfg, template_path, out_dir = scrape_complaint(complaint_id, cfg_path)
    out_name = cfg.get('output_name_pattern', 'Customer_Letter_{complaint_id}.docx').format(**values)
    out_path = out_dir / out_name
    fill_docx(str(template_path), str(out_path), values, products)
    log(f"Generated: {out_path}")
if __name__ == "__main__":
    main()